"""Build a traceable UR5 trajectory-baseline workbook from ROS 2 bags."""

from __future__ import annotations

import argparse
import base64
from dataclasses import dataclass, field
import hashlib
import json
import logging
import math
from pathlib import Path
from typing import Any, Iterable, Sequence

import numpy as np
import pandas as pd
from scipy.signal import savgol_filter

from ur5_pick_place.trajectory_metrics import ToolKinematics


LOGGER = logging.getLogger(__name__)
ARM_JOINTS = (
    'shoulder_pan_joint', 'shoulder_lift_joint', 'elbow_joint',
    'wrist_1_joint', 'wrist_2_joint', 'wrist_3_joint',
)
EVENT_TOPIC = '/trajectory_experiment/events'
POSE_TOPIC = '/world/ur5_pick_place/dynamic_pose/info'
CONTROLLER_TOPICS = (
    '/arm_controller/controller_state',
    '/joint_trajectory_controller/controller_state',
)
MIN_MOVE_RAD = 0.001
SETTLING_FRACTION = 0.02
MIN_SETTLING_BAND_RAD = 0.005
SETTLING_DWELL_S = 1.0
STEADY_WINDOW_S = 1.0
STEADY_RANGE_RAD = 0.005
STEADY_MEAN_SPEED_RAD_S = 0.01
RESAMPLE_HZ = 100.0
SAVGOL_WINDOW = 11
SAVGOL_POLYORDER = 3


@dataclass
class VectorSeries:
    times: list[float] = field(default_factory=list)
    record_times: list[float] = field(default_factory=list)
    positions: list[np.ndarray] = field(default_factory=list)
    velocities: list[np.ndarray] = field(default_factory=list)
    accelerations: list[np.ndarray] = field(default_factory=list)


@dataclass
class BagData:
    actual: VectorSeries
    reference: VectorSeries
    events: list[dict[str, Any]]
    block_poses: list[tuple[float, float, np.ndarray]]
    quality_notes: list[str]


@dataclass(frozen=True)
class StageWindow:
    name: str
    group: str
    index: int
    start: float
    endpoint: float
    end: float
    planned_duration: float


def _stamp_seconds(message: Any, bag_time_ns: int) -> float:
    stamp = getattr(getattr(message, 'header', None), 'stamp', None)
    if stamp is not None and (stamp.sec or stamp.nanosec):
        return float(stamp.sec) + float(stamp.nanosec) * 1e-9
    return bag_time_ns * 1e-9


def _point_values(point: Any, field_name: str, count: int) -> np.ndarray:
    values = getattr(point, field_name, [])
    if len(values) != count:
        return np.full(count, np.nan)
    return np.asarray(values, dtype=float)


def _named_values(names: Sequence[str], values: Sequence[float]) -> np.ndarray | None:
    if len(names) != len(values):
        return None
    lookup = dict(zip(names, values))
    if not all(name in lookup for name in ARM_JOINTS):
        return None
    result = np.asarray([lookup[name] for name in ARM_JOINTS], dtype=float)
    return result if np.all(np.isfinite(result)) else None


def _normalise(series: VectorSeries) -> VectorSeries:
    if not series.times:
        return series
    order = np.argsort(np.asarray(series.times), kind='stable')
    times = np.asarray(series.times)[order]
    keep = np.r_[np.flatnonzero(np.diff(times) != 0), len(times) - 1]

    def ordered(values: list[np.ndarray]) -> list[np.ndarray]:
        if not values:
            return []
        array = np.asarray(values)[order][keep]
        return [row for row in array]

    return VectorSeries(
        times=list(times[keep]),
        record_times=list(np.asarray(series.record_times)[order][keep]),
        positions=ordered(series.positions),
        velocities=ordered(series.velocities),
        accelerations=ordered(series.accelerations),
    )


def read_bag(path: Path) -> BagData:
    """Read feedback, controller references, experiment events and block poses."""
    try:
        import rosbag2_py
        from rclpy.serialization import deserialize_message
        from rosidl_runtime_py.utilities import get_message
    except ImportError as error:
        raise RuntimeError('Source ROS 2 Humble before reading bags') from error

    reader = rosbag2_py.SequentialReader()
    reader.open(
        rosbag2_py.StorageOptions(uri=str(path), storage_id='sqlite3'),
        rosbag2_py.ConverterOptions(
            input_serialization_format='cdr', output_serialization_format='cdr'),
    )
    available = {item.name: item.type for item in reader.get_all_topics_and_types()}
    controller_topic = next((name for name in CONTROLLER_TOPICS if name in available), None)
    if '/joint_states' not in available:
        raise ValueError(f'{path} has no /joint_states topic')
    if controller_topic is None:
        raise ValueError(f'{path} has no supported arm controller-state topic')
    selected = {'/joint_states', controller_topic, EVENT_TOPIC, POSE_TOPIC} & set(available)
    types = {name: get_message(available[name]) for name in selected}
    actual, reference = VectorSeries(), VectorSeries()
    events: list[dict[str, Any]] = []
    block_poses: list[tuple[float, float, np.ndarray]] = []
    notes: list[str] = []

    while reader.has_next():
        topic, raw, bag_ns = reader.read_next()
        if topic not in selected:
            continue
        message = deserialize_message(raw, types[topic])
        record_time = bag_ns * 1e-9
        if topic == '/joint_states':
            position = _named_values(message.name, message.position)
            if position is None:
                continue
            actual.times.append(_stamp_seconds(message, bag_ns))
            actual.record_times.append(record_time)
            actual.positions.append(position)
            velocity = _named_values(message.name, message.velocity)
            actual.velocities.append(
                velocity if velocity is not None else np.full(len(ARM_JOINTS), np.nan))
            continue
        if topic == controller_topic:
            point = getattr(message, 'reference', None)
            if point is None or len(point.positions) != len(message.joint_names):
                point = getattr(message, 'desired', None)
            if point is None:
                continue
            position = _named_values(message.joint_names, point.positions)
            if position is None:
                continue
            reference.times.append(_stamp_seconds(message, bag_ns))
            reference.record_times.append(record_time)
            reference.positions.append(position)
            reference.velocities.append(_point_values(point, 'velocities', len(ARM_JOINTS)))
            reference.accelerations.append(_point_values(point, 'accelerations', len(ARM_JOINTS)))
            continue
        if topic == EVENT_TOPIC:
            try:
                item = json.loads(message.data)
                item['record_time_s'] = record_time
                events.append(item)
            except (AttributeError, TypeError, json.JSONDecodeError):
                notes.append('Ignored malformed experiment event')
            continue
        if topic == POSE_TOPIC:
            for transform in message.transforms:
                child = transform.child_frame_id.lower()
                if 'block' not in child:
                    continue
                stamp = _stamp_seconds(transform, bag_ns)
                tr = transform.transform.translation
                block_poses.append((stamp, record_time, np.asarray([tr.x, tr.y, tr.z])))

    actual, reference = _normalise(actual), _normalise(reference)
    if len(actual.times) < 2 or len(reference.times) < 2:
        raise ValueError(f'{path} does not contain enough usable arm samples')
    if not events:
        notes.append('No stage events; stage-level results are unavailable')
    if not block_poses:
        notes.append('No Gazebo block poses; carrying status is unverified')
    return BagData(actual, reference, events, block_poses, notes)


def _median_period(times: np.ndarray) -> float:
    differences = np.diff(times)
    positive = differences[differences > 0]
    if not len(positive):
        raise ValueError('Signal timestamps do not increase')
    return float(np.median(positive))


def interpolate_with_gap_limit(
    source_times: Sequence[float], source_values: np.ndarray,
    query_times: Sequence[float], *, gap_factor: float = 3.0,
) -> tuple[np.ndarray, np.ndarray]:
    """Linearly interpolate only between nearby real samples, never extrapolate."""
    source_t = np.asarray(source_times, dtype=float)
    query_t = np.asarray(query_times, dtype=float)
    values = np.asarray(source_values, dtype=float)
    output = np.full((len(query_t), values.shape[1]), np.nan)
    valid = np.zeros(len(query_t), dtype=bool)
    limit = gap_factor * _median_period(source_t)
    right = np.searchsorted(source_t, query_t, side='left')
    for index, (time_value, after) in enumerate(zip(query_t, right)):
        if after < len(source_t) and math.isclose(source_t[after], time_value, abs_tol=1e-12):
            output[index] = values[after]
            valid[index] = np.all(np.isfinite(output[index]))
            continue
        before = after - 1
        if before < 0 or after >= len(source_t):
            continue
        gap = source_t[after] - source_t[before]
        if gap <= 0 or gap > limit:
            continue
        fraction = (time_value - source_t[before]) / gap
        output[index] = values[before] + fraction * (values[after] - values[before])
        valid[index] = np.all(np.isfinite(output[index]))
    return output, valid


def estimate_derivatives(
    actual: VectorSeries, query_times: Sequence[float],
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Estimate acceleration/jerk from measured velocity on a 100 Hz grid."""
    query_t = np.asarray(query_times, dtype=float)
    acceleration = np.full((len(query_t), len(ARM_JOINTS)), np.nan)
    jerk = np.full_like(acceleration, np.nan)
    quality = np.full(len(query_t), 'InsufficientDerivativeWindow', dtype=object)
    source_t = np.asarray(actual.times)
    velocity = np.asarray(actual.velocities)
    if len(source_t) < SAVGOL_WINDOW or not np.any(np.isfinite(velocity)):
        return acceleration, jerk, quality
    step = 1.0 / RESAMPLE_HZ
    grid = np.arange(math.ceil(source_t[0] * RESAMPLE_HZ) / RESAMPLE_HZ,
                     source_t[-1] + step * 0.1, step)
    grid_velocity, grid_valid = interpolate_with_gap_limit(source_t, velocity, grid)
    finite = grid_valid & np.all(np.isfinite(grid_velocity), axis=1)
    starts = np.flatnonzero(finite & np.r_[True, ~finite[:-1]])
    ends = np.flatnonzero(finite & np.r_[~finite[1:], True]) + 1
    grid_acc = np.full_like(grid_velocity, np.nan)
    grid_jerk = np.full_like(grid_velocity, np.nan)
    half = SAVGOL_WINDOW // 2
    for start, end in zip(starts, ends):
        if end - start < SAVGOL_WINDOW:
            continue
        segment = grid_velocity[start:end]
        grid_acc[start:end] = savgol_filter(
            segment, SAVGOL_WINDOW, SAVGOL_POLYORDER,
            deriv=1, delta=step, axis=0, mode='interp')
        grid_jerk[start:end] = savgol_filter(
            segment, SAVGOL_WINDOW, SAVGOL_POLYORDER,
            deriv=2, delta=step, axis=0, mode='interp')
        grid_acc[start:start + half] = np.nan
        grid_acc[end - half:end] = np.nan
        grid_jerk[start:start + half] = np.nan
        grid_jerk[end - half:end] = np.nan
    acceleration, acc_valid = interpolate_with_gap_limit(grid, grid_acc, query_t)
    jerk, jerk_valid = interpolate_with_gap_limit(grid, grid_jerk, query_t)
    good = acc_valid & jerk_valid
    quality[good] = 'OK'
    return acceleration, jerk, quality


def stage_windows(events: Sequence[dict[str, Any]]) -> list[StageWindow]:
    submissions = sorted(
        (event for event in events if event.get('event') == 'stage_submitted'),
        key=lambda event: float(event['sim_time_s']),
    )
    terminal_times = [float(event['sim_time_s']) for event in events
                      if event.get('event') in ('task_completed', 'task_failed')]
    windows = []
    for index, event in enumerate(submissions):
        start = float(event['sim_time_s'])
        duration = float(event.get('planned_duration_s', 0.0))
        next_start = (float(submissions[index + 1]['sim_time_s'])
                      if index + 1 < len(submissions)
                      else (max(terminal_times) if terminal_times else start + duration))
        accepted = next((
            float(candidate['sim_time_s']) for candidate in events
            if candidate.get('event') == 'stage_accepted'
            and candidate.get('stage') == event.get('stage')
            and start <= float(candidate['sim_time_s']) <= next_start
        ), start)
        windows.append(StageWindow(
            name=str(event['stage']), group=str(event.get('group', '')),
            index=int(event.get('stage_index', index + 1)), start=start,
            endpoint=accepted + max(0.0, duration),
            end=max(next_start, accepted + duration),
            planned_duration=duration,
        ))
    return windows


def _stage_for_time(time_value: float, windows: Sequence[StageWindow]) -> StageWindow | None:
    return next((window for window in windows if window.start <= time_value < window.end), None)


def aligned_joint_log(run_id: str, data: BagData, windows: Sequence[StageWindow]) -> pd.DataFrame:
    reference_t = np.asarray(data.reference.times)
    reference_q = np.asarray(data.reference.positions)
    reference_v = np.asarray(data.reference.velocities)
    actual_q, position_valid = interpolate_with_gap_limit(
        data.actual.times, np.asarray(data.actual.positions), reference_t)
    actual_v, velocity_valid = interpolate_with_gap_limit(
        data.actual.times, np.asarray(data.actual.velocities), reference_t)
    estimated_a, estimated_j, derivative_quality = estimate_derivatives(data.actual, reference_t)
    rows = []
    for sample, time_value in enumerate(reference_t):
        window = _stage_for_time(float(time_value), windows)
        for joint, name in enumerate(ARM_JOINTS):
            valid = position_valid[sample]
            rows.append({
                'Run_ID': run_id,
                'Stage': window.name if window else 'Unassigned',
                'Stage_Index': window.index if window else None,
                'Sim_Time_s': float(time_value),
                'Recording_Time_s': float(data.reference.record_times[sample]),
                'Joint': name,
                'Reference_rad': float(reference_q[sample, joint]),
                'Actual_rad': float(actual_q[sample, joint]) if valid else np.nan,
                'Error_rad': (float(reference_q[sample, joint] - actual_q[sample, joint])
                              if valid else np.nan),
                'Reference_Velocity_rad_s': (
                    float(reference_v[sample, joint])
                    if np.isfinite(reference_v[sample, joint]) else np.nan),
                'Actual_Velocity_rad_s': (
                    float(actual_v[sample, joint])
                    if velocity_valid[sample] else np.nan),
                'Estimated_Acceleration_rad_s2': (
                    float(estimated_a[sample, joint])
                    if np.isfinite(estimated_a[sample, joint]) else np.nan),
                'Estimated_Jerk_rad_s3': (
                    float(estimated_j[sample, joint])
                    if np.isfinite(estimated_j[sample, joint]) else np.nan),
                'Quality': ('OK' if valid else 'GapOrOutOfRange'),
                'Derivative_Quality': derivative_quality[sample],
            })
    return pd.DataFrame(rows)


def _time_weighted_rmse(times: np.ndarray, errors: np.ndarray) -> float:
    if len(times) < 2 or times[-1] <= times[0]:
        return np.nan
    return float(np.sqrt(np.trapz(np.square(errors), times) / (times[-1] - times[0])))


def _first_crossing(times: np.ndarray, values: np.ndarray, threshold: float, direction: float) -> float | None:
    signed = direction * (values - threshold)
    indices = np.flatnonzero(signed >= 0)
    if not len(indices):
        return None
    index = int(indices[0])
    if index == 0:
        return float(times[0])
    before, after = signed[index - 1], signed[index]
    if after == before:
        return float(times[index])
    fraction = -before / (after - before)
    return float(times[index - 1] + fraction * (times[index] - times[index - 1]))


def response_metrics(
    times: Sequence[float], reference: Sequence[float], actual: Sequence[float],
    velocities: Sequence[float], jerk: Sequence[float], *,
    endpoint_time: float, observation_end: float,
) -> dict[str, Any]:
    """Calculate one joint's trajectory and endpoint response metrics."""
    t = np.asarray(times, dtype=float)
    r = np.asarray(reference, dtype=float)
    y = np.asarray(actual, dtype=float)
    v = np.asarray(velocities, dtype=float)
    j = np.asarray(jerk, dtype=float)
    finite = np.isfinite(t) & np.isfinite(r) & np.isfinite(y)
    t, r, y, v, j = t[finite], r[finite], y[finite], v[finite], j[finite]
    if len(t) < 2:
        return {'Metric_Status': 'InsufficientSamples'}
    error = r - y
    result: dict[str, Any] = {
        'RMSE_rad': _time_weighted_rmse(t, error),
        'Max_Error_rad': float(np.max(np.abs(error))),
        'Rise_Status': 'NotEligible', 'Overshoot_Status': 'NotEligible',
        'Settling_Status': 'InsufficientObservation',
        'Steady_State_Status': 'InsufficientObservation',
        'Jerk_Status': 'InsufficientDerivativeWindow',
        'Metric_Status': 'OK',
    }
    initial, target = float(r[0]), float(r[np.argmin(np.abs(t - endpoint_time))])
    displacement = target - initial
    direction = math.copysign(1.0, displacement) if displacement else 1.0
    reference_deltas = direction * np.diff(r[t <= endpoint_time])
    monotonic = (abs(displacement) >= MIN_MOVE_RAD and
                 (not len(reference_deltas) or np.all(reference_deltas >= -1e-5)))
    if monotonic:
        low = initial + 0.1 * displacement
        high = initial + 0.9 * displacement
        t10 = _first_crossing(t, y, low, direction)
        t90 = _first_crossing(t, y, high, direction)
        if t10 is not None and t90 is not None and t90 >= t10:
            result['Trajectory_Rise_Time_s'] = t90 - t10
            result['Rise_Status'] = 'OK'
        else:
            result['Trajectory_Rise_Time_s'] = np.nan
            result['Rise_Status'] = 'ThresholdNotCrossed'
        endpoint_mask = (t >= endpoint_time) & (t <= observation_end)
        if np.any(endpoint_mask):
            excursion = direction * (y[endpoint_mask] - target)
            overshoot = max(0.0, float(np.max(excursion)))
            result['Overshoot_rad'] = overshoot
            result['Overshoot_pct'] = 100.0 * overshoot / abs(displacement)
            result['Overshoot_Status'] = 'OK'
    else:
        result['Trajectory_Rise_Time_s'] = np.nan
        result['Overshoot_rad'] = np.nan
        result['Overshoot_pct'] = np.nan

    band = max(SETTLING_FRACTION * abs(displacement), MIN_SETTLING_BAND_RAD)
    result['Settling_Band_rad'] = band
    endpoint_indices = np.flatnonzero((t >= endpoint_time) & (t <= observation_end))
    settled_time = None
    for offset, index in enumerate(endpoint_indices):
        remaining = endpoint_indices[offset:]
        if observation_end - t[index] < SETTLING_DWELL_S:
            continue
        if np.all(np.abs(y[remaining] - target) <= band):
            settled_time = float(t[index] - endpoint_time)
            break
    result['Endpoint_Settling_Time_s'] = settled_time if settled_time is not None else np.nan
    if settled_time is not None:
        result['Settling_Status'] = 'OK'
    elif endpoint_indices.size and observation_end - endpoint_time >= SETTLING_DWELL_S:
        result['Settling_Status'] = 'NotSettled'

    steady_mask = (t >= observation_end - STEADY_WINDOW_S) & (t <= observation_end)
    if (np.count_nonzero(steady_mask) >= 2 and
            t[steady_mask][-1] - t[steady_mask][0] >= 0.95 * STEADY_WINDOW_S):
        local_y = y[steady_mask]
        local_v = v[steady_mask]
        if (np.ptp(local_y) <= STEADY_RANGE_RAD and len(local_v) == len(local_y)
                and np.all(np.isfinite(local_v))
                and np.mean(np.abs(local_v)) <= STEADY_MEAN_SPEED_RAD_S):
            result['Steady_Error_rad'] = float(target - np.mean(local_y))
            result['Steady_State_Status'] = 'OK'
        else:
            result['Steady_Error_rad'] = np.nan
            result['Steady_State_Status'] = 'UnstableWindow'
    else:
        result['Steady_Error_rad'] = np.nan

    finite_jerk = j[np.isfinite(j)]
    result['Peak_Jerk_rad_s3'] = (float(np.max(np.abs(finite_jerk)))
                                  if len(finite_jerk) else np.nan)
    result['RMS_Jerk_rad_s3'] = (float(np.sqrt(np.mean(np.square(finite_jerk))))
                                 if len(finite_jerk) else np.nan)
    if len(finite_jerk):
        result['Jerk_Status'] = 'Estimated'
    endpoint_velocity = v[np.isfinite(v) & (t >= endpoint_time)]
    result['Residual_RMS_Velocity_rad_s'] = (
        float(np.sqrt(np.mean(np.square(endpoint_velocity))))
        if len(endpoint_velocity) else np.nan)
    return result


def stage_metrics(run_id: str, joint_log: pd.DataFrame,
                  windows: Sequence[StageWindow]) -> pd.DataFrame:
    rows = []
    for window in windows:
        if window.group != 'arm':
            continue
        for joint in ARM_JOINTS:
            samples = joint_log[(joint_log.Stage == window.name) &
                                (joint_log.Joint == joint) &
                                (joint_log.Quality == 'OK')]
            metrics = response_metrics(
                samples.Sim_Time_s.to_numpy(), samples.Reference_rad.to_numpy(),
                samples.Actual_rad.to_numpy(),
                samples.Actual_Velocity_rad_s.to_numpy(),
                samples.Estimated_Jerk_rad_s3.to_numpy(),
                endpoint_time=window.endpoint, observation_end=window.end,
            )
            rows.append({
                'Run_ID': run_id, 'Stage': window.name,
                'Stage_Index': window.index, 'Joint': joint,
                'Stage_Start_s': window.start, 'Reference_Endpoint_s': window.endpoint,
                'Observation_End_s': window.end, **metrics,
            })
    columns = [
        'Run_ID', 'Stage', 'Stage_Index', 'Joint', 'Stage_Start_s',
        'Reference_Endpoint_s', 'Observation_End_s', 'RMSE_rad',
        'Max_Error_rad', 'Overshoot_rad', 'Overshoot_pct',
        'Trajectory_Rise_Time_s', 'Endpoint_Settling_Time_s',
        'Settling_Band_rad', 'Steady_Error_rad', 'Peak_Jerk_rad_s3',
        'RMS_Jerk_rad_s3', 'Residual_RMS_Velocity_rad_s', 'Rise_Status',
        'Overshoot_Status', 'Settling_Status', 'Steady_State_Status',
        'Jerk_Status', 'Metric_Status',
    ]
    return pd.DataFrame(rows).reindex(columns=columns)


def load_kinematics(urdf_path: Path | None, tool_link: str) -> ToolKinematics:
    if urdf_path is not None:
        text = urdf_path.read_text()
        return ToolKinematics(text, tool_link)
    try:
        import xacro
        from ament_index_python.packages import get_package_share_directory
        source = Path(get_package_share_directory('ur5_moveit_config')) / 'config/ur5.urdf.xacro'
        document = xacro.process_file(str(source), mappings={
            'sim_ignition': 'true', 'sim_gazebo': 'false', 'name': 'IgnitionSystem',
            'simulation_controllers': str(
                source.parent / 'ros2_controllers.yaml'),
        })
        return ToolKinematics(document.toxml(), tool_link)
    except Exception as error:
        raise RuntimeError(
            'Could not build tool kinematics; supply --urdf with an expanded URDF') from error


def ee_log(run_id: str, joint_log: pd.DataFrame, kinematics: ToolKinematics) -> pd.DataFrame:
    rows = []
    good = joint_log[joint_log.Quality == 'OK']
    for (stage, time_value), sample in good.groupby(['Stage', 'Sim_Time_s'], sort=True):
        by_joint = sample.set_index('Joint')
        if not all(joint in by_joint.index for joint in ARM_JOINTS):
            continue
        desired = {joint: float(by_joint.at[joint, 'Reference_rad']) for joint in ARM_JOINTS}
        actual = {joint: float(by_joint.at[joint, 'Actual_rad']) for joint in ARM_JOINTS}
        desired_xyz = kinematics.position(desired)
        actual_xyz = kinematics.position(actual)
        rows.append({
            'Run_ID': run_id, 'Stage': stage, 'Sim_Time_s': float(time_value),
            'Desired_X_m': desired_xyz[0], 'Desired_Y_m': desired_xyz[1],
            'Desired_Z_m': desired_xyz[2], 'Actual_X_m': actual_xyz[0],
            'Actual_Y_m': actual_xyz[1], 'Actual_Z_m': actual_xyz[2],
            'Position_Error_mm': 1000.0 * math.dist(desired_xyz, actual_xyz),
            'Source': 'Calculated from joint states',
        })
    return pd.DataFrame(rows)


def carry_status(data: BagData, windows: Sequence[StageWindow],
                 kinematics: ToolKinematics) -> str:
    window = next((item for item in windows if item.name == 'move_to_place'), None)
    poses = [(time_value, xyz) for time_value, _, xyz in data.block_poses
             if window and window.start <= time_value <= window.end]
    if window is None or len(poses) < 5:
        return 'Unverified_No_Block_Pose'
    times = np.asarray([item[0] for item in poses])
    joints, valid = interpolate_with_gap_limit(
        data.actual.times, np.asarray(data.actual.positions), times)
    relative = []
    block = []
    for joint_values, (_, xyz), usable in zip(joints, poses, valid):
        if not usable:
            continue
        tool = kinematics.position(dict(zip(ARM_JOINTS, joint_values)))
        relative.append(xyz - np.asarray(tool))
        block.append(xyz)
    if len(relative) < 5:
        return 'Unverified_Insufficient_Pose_Alignment'
    travel = math.dist(block[0], block[-1])
    spread = np.max(np.linalg.norm(relative - np.median(relative, axis=0), axis=1))
    if travel < 0.05:
        return 'Not_Carried'
    return 'Verified_Carried' if spread <= 0.02 else 'Unverified_Unstable_Relative_Pose'


def load_planned_path(path: Path) -> tuple[str, pd.DataFrame]:
    """Decode the exact saved RobotTrajectory bundle into long-form waypoints."""
    try:
        from moveit_msgs.msg import RobotTrajectory
        from rclpy.serialization import deserialize_message
    except ImportError as error:
        raise RuntimeError('Source ROS 2 Humble before decoding a trajectory bundle') from error
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    payload = json.loads(path.read_text())
    rows = []
    for stage_index, item in enumerate(payload['stages'], start=1):
        trajectory = deserialize_message(
            base64.b64decode(item['trajectory']), RobotTrajectory).joint_trajectory
        for waypoint, point in enumerate(trajectory.points):
            time_value = point.time_from_start.sec + point.time_from_start.nanosec * 1e-9
            for joint_index, joint in enumerate(trajectory.joint_names):
                rows.append({
                    'Stage': item['name'], 'Stage_Index': stage_index,
                    'Waypoint_Index': waypoint, 'Joint': joint,
                    'Time_From_Start_s': time_value,
                    'Position_rad': point.positions[joint_index],
                    'Velocity_rad_s': (point.velocities[joint_index]
                                       if len(point.velocities) == len(trajectory.joint_names)
                                       else np.nan),
                    'Acceleration_rad_s2': (point.accelerations[joint_index]
                                            if len(point.accelerations) == len(trajectory.joint_names)
                                            else np.nan),
                })
    return digest, pd.DataFrame(rows)


def definitions_frame() -> pd.DataFrame:
    rows = [
        ('Tracking RMSE', 'rad', 'sqrt(time integral(error^2) / duration)', 'At least two aligned samples'),
        ('Maximum tracking error', 'rad', 'max(abs(reference - actual))', 'Aligned samples'),
        ('End-effector tracking error', 'mm', 'Euclidean FK position difference', 'Six valid arm joints'),
        ('Maximum endpoint overshoot', 'rad, %', 'Motion-direction excursion beyond endpoint', 'Monotonic move >= 0.001 rad'),
        ('Trajectory rise time', 's', 'First 10% to first 90% actual crossing', 'Monotonic move >= 0.001 rad'),
        ('Endpoint settling time', 's', 'First entry that remains in band through observation end', 'At least 1 s observation'),
        ('Settling band', 'rad', 'max(2% of move, 0.005 rad)', 'All arm stages'),
        ('Steady-state error', 'rad', 'target - mean(actual) over final 1 s', 'range <= 0.005 rad and mean speed <= 0.01 rad/s'),
        ('Estimated jerk', 'rad/s^3', '2nd derivative of measured velocity', '100 Hz, Savitzky-Golay window 11, order 3'),
        ('Residual motion', 'rad/s', 'RMS measured velocity after reference endpoint', 'Available endpoint samples'),
        ('Carry status', 'category', 'Block travel and tool-relative pose stability', 'Gazebo block pose bridge available'),
        ('Joint torque', 'N m', 'Unavailable for this baseline', 'Joint-state effort values must be finite'),
    ]
    return pd.DataFrame(rows, columns=['Measurement', 'Units', 'Formula', 'Eligibility_or_Threshold'])


def _event_config(path: Path, data: BagData, defaults: dict[str, Any]) -> dict[str, Any]:
    event = next((item for item in data.events if item.get('event') == 'task_started'), {})
    return {
        'Run_ID': event.get('run_id') or path.name,
        'Controller_Mode': event.get('controller_mode', defaults['controller_mode']),
        'Trajectory_ID': event.get('trajectory_id', defaults['trajectory_id']),
        'Block_Mass_kg': float(event.get('block_mass_kg', defaults['block_mass_kg'])),
        'Payload_Condition': event.get('payload_condition', defaults['payload_condition']),
        'Velocity_Scaling': float(event.get('arm_velocity_scaling', defaults['velocity_scaling'])),
        'Acceleration_Scaling': float(event.get('arm_acceleration_scaling', defaults['acceleration_scaling'])),
        'Frame': defaults['frame'], 'Tool_Link': defaults['tool_link'],
        'Bag_Path': str(path.resolve()),
    }


def analyse_baseline_run(path: Path, trajectory_id: str, kinematics: ToolKinematics,
                         defaults: dict[str, Any]) -> dict[str, pd.DataFrame | dict[str, Any]]:
    data = read_bag(path)
    windows = stage_windows(data.events)
    config = _event_config(path, data, {**defaults, 'trajectory_id': trajectory_id})
    config['Carry_Status'] = carry_status(data, windows, kinematics)
    config['Data_Quality_Notes'] = '; '.join(data.quality_notes) or 'OK'
    joint = aligned_joint_log(config['Run_ID'], data, windows)
    ee = ee_log(config['Run_ID'], joint, kinematics)
    metrics = stage_metrics(config['Run_ID'], joint, windows)
    terminal = next((item for item in reversed(data.events)
                     if item.get('event') in ('task_completed', 'task_failed')), None)
    started = next((item for item in data.events if item.get('event') == 'task_started'), None)
    execution_status = ('Success' if terminal and terminal.get('event') == 'task_completed'
                        else ('Failed' if terminal else 'Unknown'))
    duration = (float(terminal['sim_time_s']) - float(started['sim_time_s'])
                if terminal and started else data.actual.times[-1] - data.actual.times[0])
    overall_rmse = []
    for _, samples in joint[joint.Quality == 'OK'].groupby('Joint'):
        overall_rmse.append(_time_weighted_rmse(
            samples.Sim_Time_s.to_numpy(), samples.Error_rad.to_numpy()))
    summary = {
        'Run_ID': config['Run_ID'], 'Controller_Mode': config['Controller_Mode'],
        'Trajectory_ID': config['Trajectory_ID'],
        'Payload_Condition': config['Payload_Condition'],
        'Execution_Status': execution_status, 'Task_Duration_s': duration,
        'Max_Joint_RMSE_rad': (float(np.nanmax(overall_rmse))
                               if overall_rmse else np.nan),
        'EE_RMSE_mm': (float(np.sqrt(np.mean(np.square(ee.Position_Error_mm))))
                       if not ee.empty else np.nan),
        'EE_Max_Error_mm': (float(ee.Position_Error_mm.max()) if not ee.empty else np.nan),
        'Valid_Stage_Count': (int(metrics[metrics.Metric_Status == 'OK'].Stage.nunique())
                              if not metrics.empty else 0),
        'Quality_Status': 'OK' if not data.quality_notes else '; '.join(data.quality_notes),
    }
    controller = pd.DataFrame([{
        'Run_ID': config['Run_ID'], 'Joint': joint_name,
        'Command_Interface': 'position', 'Open_Loop': True,
        'Kp': np.nan, 'Ki': np.nan, 'Kd': np.nan,
        'Gain_Status': 'Unavailable: position-forwarding open-loop configuration',
    } for joint_name in ARM_JOINTS])
    return {'config': config, 'controller': controller, 'joint': joint,
            'ee': ee, 'metrics': metrics, 'summary': summary}


def statistical_summary(metrics: pd.DataFrame, config: pd.DataFrame) -> pd.DataFrame:
    if metrics.empty:
        return pd.DataFrame(columns=[
            'Controller_Mode', 'Trajectory_ID', 'Payload_Condition', 'Stage',
            'Joint', 'Metric', 'Valid_N', 'Mean', 'Standard_Deviation'])
    value_columns = [
        'RMSE_rad', 'Max_Error_rad', 'Overshoot_rad', 'Overshoot_pct',
        'Trajectory_Rise_Time_s', 'Endpoint_Settling_Time_s', 'Steady_Error_rad',
        'Peak_Jerk_rad_s3', 'RMS_Jerk_rad_s3', 'Residual_RMS_Velocity_rad_s',
    ]
    merged = metrics.merge(config[[
        'Run_ID', 'Controller_Mode', 'Trajectory_ID', 'Payload_Condition']], on='Run_ID')
    long = merged.melt(
        id_vars=['Controller_Mode', 'Trajectory_ID', 'Payload_Condition', 'Stage', 'Joint'],
        value_vars=value_columns, var_name='Metric', value_name='Value').dropna(subset=['Value'])
    group = ['Controller_Mode', 'Trajectory_ID', 'Payload_Condition', 'Stage', 'Joint', 'Metric']
    result = long.groupby(group, dropna=False).Value.agg(['count', 'mean', 'std']).reset_index()
    return result.rename(columns={
        'count': 'Valid_N', 'mean': 'Mean', 'std': 'Standard_Deviation'})


def _format_workbook(output_path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.chart import LineChart, Reference, ScatterChart, Series
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = load_workbook(output_path)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 28
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F4E78')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(42, max(11, max(len(str(cell.value or '')) for cell in column[:200]) + 2))
            sheet.column_dimensions[letter].width = width

    charts = workbook['Charts']
    data = workbook['_Chart_Data']
    if data.max_row > 2:
        joint_end = data['J1'].value or 1
        position = LineChart()
        position.title = 'First run: desired and actual joint position'
        position.y_axis.title = 'Position (rad)'
        position.x_axis.title = 'Simulation time (s)'
        for column in (2, 3):
            position.add_data(Reference(data, min_col=column, min_row=1, max_row=joint_end),
                              titles_from_data=True, from_rows=False)
        position.set_categories(Reference(data, min_col=1, min_row=2, max_row=joint_end))
        position.height, position.width = 8, 15
        charts.add_chart(position, 'A3')

        tracking_error = LineChart()
        tracking_error.title = 'First run: joint error and settling band'
        tracking_error.y_axis.title = 'Error (rad)'
        tracking_error.x_axis.title = 'Simulation time (s)'
        for column in (4, 8, 9):
            tracking_error.add_data(
                Reference(data, min_col=column, min_row=1, max_row=joint_end),
                titles_from_data=True)
        tracking_error.set_categories(
            Reference(data, min_col=1, min_row=2, max_row=joint_end))
        tracking_error.height, tracking_error.width = 8, 15
        charts.add_chart(tracking_error, 'A20')

        smooth = LineChart()
        smooth.title = 'First run: velocity, estimated acceleration and jerk'
        for column in (5, 6, 7):
            smooth.add_data(Reference(data, min_col=column, min_row=1, max_row=joint_end),
                            titles_from_data=True)
        smooth.set_categories(Reference(data, min_col=1, min_row=2, max_row=joint_end))
        smooth.height, smooth.width = 8, 15
        charts.add_chart(smooth, 'A37')

        ee_end = data['J2'].value or 1
        xy = ScatterChart()
        xy.title = 'First run: desired and actual tool path (XY)'
        xy.x_axis.title, xy.y_axis.title = 'X (m)', 'Y (m)'
        for x_col, y_col, title in ((11, 12, 'Desired XY'), (14, 15, 'Actual XY')):
            series = Series(Reference(data, min_col=y_col, min_row=2, max_row=ee_end),
                            Reference(data, min_col=x_col, min_row=2, max_row=ee_end),
                            title=title)
            xy.series.append(series)
        xy.height, xy.width = 8, 15
        charts.add_chart(xy, 'Q3')
        xz = ScatterChart()
        xz.title = 'First run: desired and actual tool path (XZ)'
        xz.x_axis.title, xz.y_axis.title = 'X (m)', 'Z (m)'
        for x_col, z_col, title in ((11, 13, 'Desired XZ'), (14, 16, 'Actual XZ')):
            xz.series.append(Series(
                Reference(data, min_col=z_col, min_row=2, max_row=ee_end),
                Reference(data, min_col=x_col, min_row=2, max_row=ee_end), title=title))
        xz.height, xz.width = 8, 15
        charts.add_chart(xz, 'Q20')

        stats_end = data['J3'].value or 1
        error = LineChart()
        error.title = 'Per-stage RMSE: mean and standard deviation'
        error.add_data(Reference(data, min_col=22, max_col=23, min_row=1,
                                 max_row=stats_end), titles_from_data=True)
        error.set_categories(Reference(data, min_col=20, min_row=2, max_row=stats_end))
        error.height, error.width = 8, 15
        charts.add_chart(error, 'Q37')
    data.sheet_state = 'hidden'
    workbook.save(output_path)


def export_workbook(results: Sequence[dict[str, Any]], planned: pd.DataFrame,
                    output_path: Path) -> None:
    configs = pd.DataFrame([item['config'] for item in results])
    controllers = pd.concat([item['controller'] for item in results], ignore_index=True)
    joints = pd.concat([item['joint'] for item in results], ignore_index=True)
    ee = pd.concat([item['ee'] for item in results], ignore_index=True)
    metrics = pd.concat([item['metrics'] for item in results], ignore_index=True)
    summaries = pd.DataFrame([item['summary'] for item in results])
    summaries['Successful_Trials_Total'] = int(
        (summaries.Execution_Status == 'Success').sum())
    summaries['Recorded_Trials_Total'] = len(summaries)
    stats = statistical_summary(metrics, configs)
    planned_runs = pd.concat([
        planned.assign(Run_ID=item['config']['Run_ID']) for item in results], ignore_index=True)
    planned_runs = planned_runs[['Run_ID', *[column for column in planned.columns]]]
    first_run = configs.iloc[0].Run_ID
    chart_joint = joints[(joints.Run_ID == first_run) &
                         (joints.Joint == ARM_JOINTS[0])].copy()
    if metrics.empty:
        chart_joint['Positive_Settling_Band_rad'] = np.nan
        chart_joint['Negative_Settling_Band_rad'] = np.nan
    else:
        band_lookup = metrics[(metrics.Run_ID == first_run) &
                              (metrics.Joint == ARM_JOINTS[0])].set_index(
                                  'Stage').Settling_Band_rad
        chart_joint['Positive_Settling_Band_rad'] = chart_joint.Stage.map(band_lookup)
        chart_joint['Negative_Settling_Band_rad'] = -chart_joint.Stage.map(band_lookup)
    chart_joint = chart_joint.rename(columns={
        'Actual_Velocity_rad_s': 'Velocity_rad_s',
        'Estimated_Acceleration_rad_s2': 'Acceleration_Estimate_rad_s2',
        'Estimated_Jerk_rad_s3': 'Jerk_Estimate_rad_s3',
    })[[
        'Sim_Time_s', 'Reference_rad', 'Actual_rad', 'Error_rad', 'Velocity_rad_s',
        'Acceleration_Estimate_rad_s2', 'Jerk_Estimate_rad_s3',
        'Positive_Settling_Band_rad', 'Negative_Settling_Band_rad',
    ]]
    chart_ee = ee[ee.Run_ID == first_run][[
        'Desired_X_m', 'Desired_Y_m', 'Desired_Z_m',
        'Actual_X_m', 'Actual_Y_m', 'Actual_Z_m']]
    chart_stats = stats[(stats.Joint == ARM_JOINTS[0]) &
                        (stats.Metric == 'RMSE_rad')][[
                            'Stage', 'Metric', 'Mean', 'Standard_Deviation']]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        definitions_frame().to_excel(writer, sheet_name='Definitions', index=False)
        configs.to_excel(writer, sheet_name='Run_Config', index=False)
        controllers.to_excel(writer, sheet_name='Controller_Config', index=False)
        planned_runs.to_excel(writer, sheet_name='Planned_Path', index=False)
        joints.to_excel(writer, sheet_name='Joint_Log', index=False)
        ee.to_excel(writer, sheet_name='EE_Log', index=False)
        metrics.to_excel(writer, sheet_name='Stage_Metrics', index=False)
        summaries.to_excel(writer, sheet_name='Run_Summary', index=False)
        stats.to_excel(writer, sheet_name='Statistical_Summary', index=False)
        pd.DataFrame({'Dashboard': ['Charts use the first run; filter data sheets for all trials.']}).to_excel(
            writer, sheet_name='Charts', index=False)
        chart_joint.to_excel(writer, sheet_name='_Chart_Data', index=False, startcol=0)
        chart_ee.to_excel(writer, sheet_name='_Chart_Data', index=False, startcol=10)
        chart_stats.to_excel(writer, sheet_name='_Chart_Data', index=False, startcol=19)
        worksheet = writer.sheets['_Chart_Data']
        worksheet['J1'] = len(chart_joint) + 1
        worksheet['J2'] = len(chart_ee) + 1
        worksheet['J3'] = len(chart_stats) + 1
    _format_workbook(output_path)


def default_trajectory_path() -> Path:
    try:
        from ament_index_python.packages import get_package_share_directory
        installed = Path(get_package_share_directory('ur5_pick_place')) / 'baseline/selected_trajectory.json'
        if installed.exists():
            return installed
    except Exception:
        pass
    return Path('experiment_bags/Optimizer_0kg_run04/selected_trajectory.json')


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input_dir', type=Path, help='Directory containing trial bag directories')
    parser.add_argument('-o', '--output', type=Path, default=Path('UR5_Baseline_Results.xlsx'))
    parser.add_argument('--trajectory', type=Path, default=default_trajectory_path())
    parser.add_argument('--urdf', type=Path, help='Expanded URDF; default processes workspace xacro')
    parser.add_argument('--controller-mode', default='Position_OpenLoop')
    parser.add_argument('--payload-condition', default='block_0.1kg_contact_grasp')
    parser.add_argument('--block-mass-kg', type=float, default=0.1)
    parser.add_argument('--velocity-scaling', type=float, default=0.2)
    parser.add_argument('--acceleration-scaling', type=float, default=0.2)
    parser.add_argument('--frame', default='world')
    parser.add_argument('--tool-link', default='tool0')
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')
    args = build_parser().parse_args(argv)
    if not args.input_dir.is_dir():
        raise SystemExit(f'Input directory does not exist: {args.input_dir}')
    if not args.trajectory.is_file():
        raise SystemExit(f'Trajectory bundle does not exist: {args.trajectory}')
    digest, planned = load_planned_path(args.trajectory)
    trajectory_id = f'run04-{digest[:12]}'
    kinematics = load_kinematics(args.urdf, args.tool_link)
    defaults = vars(args)
    bags = sorted(path for path in args.input_dir.iterdir()
                  if path.is_dir() and (path / 'metadata.yaml').exists())
    results = []
    for path in bags:
        try:
            results.append(analyse_baseline_run(path, trajectory_id, kinematics, defaults))
        except (OSError, RuntimeError, ValueError, KeyError) as error:
            LOGGER.error('Skipping %s: %s', path.name, error)
    if not results:
        raise SystemExit('No usable trial bags were found')
    export_workbook(results, planned, args.output)
    LOGGER.info('Wrote %d trial(s) to %s', len(results), args.output)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
