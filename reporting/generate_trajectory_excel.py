#!/usr/bin/env python3
"""Generate a read-only UR5 trajectory-analysis Excel workbook from a CSV log.

This is an offline reporting utility.  It does not import ROS, contact ROS nodes,
or modify the input CSV.  All metrics use the full-resolution source data.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import math
import os
import shutil
import statistics
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Sequence

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.table import Table, TableStyleInfo


JOINTS = (
    ("shoulder_pan_joint", "J1", "Shoulder Pan"),
    ("shoulder_lift_joint", "J2", "Shoulder Lift"),
    ("elbow_joint", "J3", "Elbow"),
    ("wrist_1_joint", "J4", "Wrist 1"),
    ("wrist_2_joint", "J5", "Wrist 2"),
    ("wrist_3_joint", "J6", "Wrist 3"),
)

REQUIRED_BASE = ("time_sec", "trajectory_id", "controller_mode", "payload_kg")
EE_COLUMNS = (
    "reference_ee_x_m", "reference_ee_y_m", "reference_ee_z_m",
    "actual_ee_x_m", "actual_ee_y_m", "actual_ee_z_m",
    "ee_error_x_mm", "ee_error_y_mm", "ee_error_z_mm", "ee_error_3d_mm",
)

NAVY = "17365D"
BLUE = "1F5BFF"
RED = "E31A1C"
GREEN = "15931B"
ORANGE = "F39C12"
PURPLE = "6C4CCF"
BLACK = "111111"
WHITE = "FFFFFF"
PALE_BLUE = "DDEEFF"
PALE_GREEN = "DFF2DA"
PALE_ORANGE = "FFF0CD"
PALE_PURPLE = "EDE2FA"
PALE_RED = "F8D8E2"
GRID = "D9E1EA"
TEXT = "14213D"
THIN = Side(style="thin", color="C9D3E1")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="Source joint_trajectory_log.csv")
    parser.add_argument("--output", required=True, type=Path, help="New .xlsx report path")
    parser.add_argument("--overwrite", action="store_true", help="Allow replacing an existing output report")
    return parser.parse_args()


def finite(value) -> bool:
    return isinstance(value, (int, float)) and math.isfinite(value)


def as_number(text: str):
    if text == "":
        return None
    stripped = text.strip()
    if stripped.lower() in {"nan", "+nan", "-nan", "inf", "+inf", "-inf", "infinity", "+infinity", "-infinity"}:
        return None
    try:
        return float(stripped)
    except ValueError:
        return text


def valid(values: Iterable) -> list[float]:
    return [float(v) for v in values if finite(v)]


def mean(values: Iterable) -> float | None:
    vals = valid(values)
    return sum(vals) / len(vals) if vals else None


def rmse(values: Iterable) -> float | None:
    vals = valid(values)
    return math.sqrt(sum(v * v for v in vals) / len(vals)) if vals else None


def max_abs(values: Iterable) -> float | None:
    vals = valid(values)
    return max(abs(v) for v in vals) if vals else None


def peak_abs_info(values: Sequence, times: Sequence):
    """Return (signed value, absolute value, relative time) for the largest valid magnitude."""
    candidates = [(abs(float(v)), float(v), times[i]) for i, v in enumerate(values) if finite(v) and finite(times[i])]
    if not candidates:
        return None, None, None
    _, signed, time = max(candidates, key=lambda item: item[0])
    return signed, abs(signed), float(time)


def pop_std(values: Iterable) -> float | None:
    vals = valid(values)
    return statistics.pstdev(vals) if vals else None


def maximum(values: Iterable) -> float | None:
    vals = valid(values)
    return max(vals) if vals else None


def last_valid(values: Sequence) -> float | None:
    for value in reversed(values):
        if finite(value):
            return float(value)
    return None


def fmt(value, digits=3, suffix="") -> str:
    if value is None or not finite(value):
        return "N/A"
    return f"{value:.{digits}f}{suffix}"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_csv(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        try:
            headers = next(reader)
        except StopIteration as exc:
            raise ValueError("Input CSV is empty") from exc
        raw_rows = list(reader)
    if not raw_rows:
        raise ValueError("Input CSV has a header but no samples")
    if any(len(row) != len(headers) for row in raw_rows):
        bad = [i + 2 for i, row in enumerate(raw_rows) if len(row) != len(headers)]
        raise ValueError(f"CSV row width differs from header at rows {bad[:10]}")
    missing = [name for name in REQUIRED_BASE if name not in headers]
    for joint, _, _ in JOINTS:
        for suffix in (
            "reference_position_rad", "actual_position_rad", "position_error_rad",
            "reference_velocity_rad_s", "actual_velocity_rad_s",
        ):
            name = f"{joint}__{suffix}"
            if name not in headers:
                missing.append(name)
    missing.extend(name for name in EE_COLUMNS if name not in headers)
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(missing))
    values = [[as_number(cell) for cell in row] for row in raw_rows]
    by_name = {name: [row[i] for row in values] for i, name in enumerate(headers)}
    return headers, raw_rows, values, by_name


def unique_text(values: Sequence) -> str:
    result = []
    for value in values:
        if value is not None and str(value).strip() and str(value) not in result:
            result.append(str(value))
    return ", ".join(result) if result else "Not recorded"


def calculate(by_name: dict[str, list], row_count: int):
    times = by_name["time_sec"]
    first_time = next((float(v) for v in times if finite(v)), None)
    rel_time = [float(v) - first_time if finite(v) and first_time is not None else None for v in times]
    valid_times = valid(times)
    duration = valid_times[-1] - valid_times[0] if len(valid_times) >= 2 else None
    intervals = [b - a for a, b in zip(valid_times, valid_times[1:]) if b - a > 0]
    median_dt = statistics.median(intervals) if intervals else None
    nonpositive = sum(1 for a, b in zip(valid_times, valid_times[1:]) if b - a <= 0)

    calc_headers = ["relative_time_s"]
    calc_rows = [[rel_time[i]] for i in range(row_count)]
    joint_arrays = {}
    discrepancy_rows = []
    activity_speed = [None] * row_count
    for joint, jid, label in JOINTS:
        ref = by_name[f"{joint}__reference_position_rad"]
        actual = by_name[f"{joint}__actual_position_rad"]
        logged = by_name[f"{joint}__position_error_rad"]
        vel = by_name[f"{joint}__actual_velocity_rad_s"]
        ref_deg = [math.degrees(v) if finite(v) else None for v in ref]
        actual_deg = [math.degrees(v) if finite(v) else None for v in actual]
        recomputed_rad = [r - a if finite(r) and finite(a) else None for r, a in zip(ref, actual)]
        differences = [abs(l - c) for l, c in zip(logged, recomputed_rad) if finite(l) and finite(c)]
        mismatch = sum(1 for d in differences if d > 1e-9)
        use_logged = bool(valid(logged)) and mismatch == 0
        selected_rad = logged if use_logged else recomputed_rad
        error_deg = [math.degrees(v) if finite(v) else None for v in selected_rad]
        velocity_deg = [math.degrees(v) if finite(v) else None for v in vel]
        for i, value in enumerate(velocity_deg):
            if finite(value):
                activity_speed[i] = max(activity_speed[i] or 0.0, abs(value))
        joint_arrays[jid] = {
            "name": label, "ref": ref_deg, "actual": actual_deg,
            "error": error_deg, "velocity": velocity_deg,
            "error_source": "logged" if use_logged else "recomputed reference minus actual",
        }
        discrepancy_rows.append((jid, len(differences), mismatch, maximum(differences), use_logged))
        new_headers = [f"{jid}_reference_deg", f"{jid}_actual_deg", f"{jid}_error_deg", f"{jid}_actual_velocity_deg_s"]
        calc_headers.extend(new_headers)
        for i in range(row_count):
            calc_rows[i].extend([ref_deg[i], actual_deg[i], error_deg[i], velocity_deg[i]])

    ee_arrays = {name: by_name[name] for name in EE_COLUMNS}
    ee_component_check = []
    for axis in "xyz":
        ref = by_name[f"reference_ee_{axis}_m"]
        actual = by_name[f"actual_ee_{axis}_m"]
        logged = by_name[f"ee_error_{axis}_mm"]
        recomputed = [(r - a) * 1000.0 if finite(r) and finite(a) else None for r, a in zip(ref, actual)]
        differences = [abs(l - c) for l, c in zip(logged, recomputed) if finite(l) and finite(c)]
        ee_component_check.append((axis.upper(), len(differences), sum(d > 1e-6 for d in differences), maximum(differences)))
    recomputed_3d = []
    for x, y, z in zip(by_name["ee_error_x_mm"], by_name["ee_error_y_mm"], by_name["ee_error_z_mm"]):
        recomputed_3d.append(math.sqrt(x*x + y*y + z*z) if finite(x) and finite(y) and finite(z) else None)
    mag_diff = [abs(a-b) for a, b in zip(by_name["ee_error_3d_mm"], recomputed_3d) if finite(a) and finite(b)]
    ee_magnitude_check = (len(mag_diff), sum(d > 1e-6 for d in mag_diff), maximum(mag_diff))

    ee_calc_names = list(EE_COLUMNS)
    calc_headers.extend(ee_calc_names)
    calc_headers.append("zero_reference")
    for i in range(row_count):
        calc_rows[i].extend([by_name[name][i] for name in ee_calc_names])
        calc_rows[i].append(0.0 if finite(rel_time[i]) else None)

    activity_threshold = 1.0
    moving_mask = [finite(v) and v >= activity_threshold for v in activity_speed]
    moving_indices = [i for i, flag in enumerate(moving_mask) if flag]
    activity_start = rel_time[moving_indices[0]] if moving_indices else None
    activity_end = rel_time[moving_indices[-1]] if moving_indices else None

    ee_error = by_name["ee_error_3d_mm"]
    peak = maximum(ee_error)
    peak_index = next((i for i, v in enumerate(ee_error) if finite(v) and v == peak), None)
    peak_time = rel_time[peak_index] if peak_index is not None else None
    overall_ee = {
        "rmse": rmse(ee_error), "max": peak, "final": last_valid(ee_error),
        "mean": mean(ee_error), "std": pop_std(ee_error), "peak_time": peak_time,
        "valid": len(valid(ee_error)),
    }
    moving_ee_values = [v for v, flag in zip(ee_error, moving_mask) if flag]
    moving_ee = {
        "rmse": rmse(moving_ee_values), "max": maximum(moving_ee_values),
        "mean": mean(moving_ee_values), "std": pop_std(moving_ee_values),
        "valid": len(valid(moving_ee_values)),
    }
    return {
        "first_time": first_time, "last_time": valid_times[-1] if valid_times else None,
        "duration": duration, "median_dt": median_dt, "nonpositive_intervals": nonpositive,
        "rel_time": rel_time, "calc_headers": calc_headers, "calc_rows": calc_rows,
        "joint_arrays": joint_arrays, "joint_discrepancy": discrepancy_rows,
        "ee_component_check": ee_component_check, "ee_magnitude_check": ee_magnitude_check,
        "activity_threshold": activity_threshold, "moving_mask": moving_mask,
        "activity_start": activity_start, "activity_end": activity_end,
        "moving_count": len(moving_indices), "overall_ee": overall_ee, "moving_ee": moving_ee,
    }


def set_cell(ws, row, col, value, *, fill=None, font=None, align=None, border=None, number_format=None):
    cell = ws.cell(row, col, value)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if font:
        cell.font = font
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    if number_format:
        cell.number_format = number_format
    return cell


def merged_label(ws, cell_range, text, fill, font_size=12, color=TEXT, horizontal="left"):
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = text
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Aptos", size=font_size, bold=True, color=color)
    cell.alignment = Alignment(horizontal=horizontal, vertical="center")
    min_col = ws[cell_range.split(":")[0]].column
    max_col = ws[cell_range.split(":")[1]].column
    row = ws[cell_range.split(":")[0]].row
    for col in range(min_col, max_col + 1):
        ws.cell(row, col).fill = PatternFill("solid", fgColor=fill)


def style_sheet(ws, gridlines=True, zoom=85):
    ws.sheet_view.showGridLines = gridlines
    ws.sheet_view.zoomScale = zoom
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.outlinePr.summaryBelow = True


def write_raw_data(wb, headers, raw_rows):
    ws = wb.create_sheet("Raw_Data")
    style_sheet(ws, False, 75)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(raw_rows)+1}"
    for col, name in enumerate(headers, 1):
        set_cell(ws, 1, col, name, fill=NAVY, font=Font(name="Aptos", size=9, bold=True, color=WHITE),
                 align=Alignment(horizontal="center", vertical="center", wrap_text=True), border=Border(bottom=THIN))
        ws.column_dimensions[get_column_letter(col)].width = 17 if col > 4 else (21 if col in (2,3) else 14)
    ws.row_dimensions[1].height = 45
    for r, row in enumerate(raw_rows, 2):
        for c, text in enumerate(row, 1):
            value = as_number(text)
            cell = ws.cell(r, c, value)
            if isinstance(value, float):
                cell.number_format = "0.000000000000000"
        if r % 2 == 0:
            for cell in ws[r]:
                cell.fill = PatternFill("solid", fgColor="F7F9FC")
    return ws


def write_calculated(wb, calc):
    ws = wb.create_sheet("Calculated_Data")
    style_sheet(ws, False, 80)
    ws.freeze_panes = "A4"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(calc["calc_headers"]))
    set_cell(ws, 1, 1, "Calculated chart arrays — values derived reproducibly from Raw_Data",
             fill=NAVY, font=Font(name="Aptos Display", size=15, bold=True, color=WHITE),
             align=Alignment(horizontal="left", vertical="center"))
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(calc["calc_headers"]))
    set_cell(ws, 2, 1,
             "Method: relative time = time_sec − first valid time; radians × 180/π; error uses validated logged reference-minus-actual values, otherwise recomputation; no interpolation.",
             fill=PALE_BLUE, font=Font(name="Aptos", size=10, color=TEXT), align=Alignment(wrap_text=True))
    for c, name in enumerate(calc["calc_headers"], 1):
        set_cell(ws, 3, c, name, fill="4472C4", font=Font(name="Aptos", size=9, bold=True, color=WHITE),
                 align=Alignment(horizontal="center", vertical="center", wrap_text=True))
        ws.column_dimensions[get_column_letter(c)].width = max(15, min(25, len(name) + 2))
    ws.row_dimensions[3].height = 38
    for r, row in enumerate(calc["calc_rows"], 4):
        for c, value in enumerate(row, 1):
            cell = ws.cell(r, c, value if finite(value) else None)
            cell.number_format = "0.000000"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(calc['calc_headers']))}{len(calc['calc_rows'])+3}"
    return ws


def title_sheet(ws, title, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    set_cell(ws, 1, 1, title, fill=NAVY, font=Font(name="Aptos Display", size=16, bold=True, color=WHITE),
             align=Alignment(horizontal="left", vertical="center"))
    ws.row_dimensions[1].height = 28


def availability(by_name, patterns: Sequence[str]) -> str:
    columns = [name for name in by_name if any(pattern in name for pattern in patterns)]
    count = sum(len(valid(by_name[name])) for name in columns)
    return f"Recorded ({count:,} valid cells)" if count else "Not recorded"


def write_metrics(wb, by_name, calc):
    ws = wb.create_sheet("Metrics")
    style_sheet(ws, False, 90)
    title_sheet(ws, "UR5 Trajectory Analysis — Metrics, Definitions, and Validity", 9)
    headers = ["Scope", "Metric", "Value", "Unit", "Valid samples", "Definition / method", "Status", "Notes", "Source"]
    for c, name in enumerate(headers, 1):
        set_cell(ws, 3, c, name, fill="4472C4", font=Font(name="Aptos", bold=True, color=WHITE),
                 align=Alignment(horizontal="center", vertical="center", wrap_text=True), border=Border(bottom=THIN))
    rows = []
    n = len(by_name["time_sec"])
    for _, jid, label in JOINTS:
        errors = calc["joint_arrays"][jid]["error"]
        moving = [v for v, flag in zip(errors, calc["moving_mask"]) if flag]
        rows.extend([
            ("Overall", f"{jid} {label} position RMSE", rmse(errors), "deg", len(valid(errors)), "sqrt(mean(signed_error²))", "Valid" if valid(errors) else "N/A", calc["joint_arrays"][jid]["error_source"], "Calculated_Data"),
            ("Overall", f"{jid} {label} maximum absolute position error", max_abs(errors), "deg", len(valid(errors)), "max(abs(signed_error))", "Valid" if valid(errors) else "N/A", "", "Calculated_Data"),
            ("Moving samples", f"{jid} {label} position RMSE", rmse(moving), "deg", len(valid(moving)), "same RMSE, only samples meeting activity threshold", "Valid" if valid(moving) else "N/A", f"max |actual velocity| across joints ≥ {calc['activity_threshold']:.1f} deg/s", "Calculated_Data"),
        ])
    ee = calc["overall_ee"]
    moving_ee = calc["moving_ee"]
    rows.extend([
        ("Overall", "EE 3D position RMSE", ee["rmse"], "mm", ee["valid"], "sqrt(mean(ee_error_3d_mm²))", "Valid" if ee["valid"] else "N/A", "Translational error magnitude only; no orientation tracking", "Raw_Data"),
        ("Overall", "EE maximum 3D position error", ee["max"], "mm", ee["valid"], "max(ee_error_3d_mm)", "Valid" if ee["valid"] else "N/A", f"Peak at relative time {fmt(ee['peak_time'],3,' s')}", "Raw_Data"),
        ("Overall", "EE final logged 3D position error", ee["final"], "mm", ee["valid"], "last valid logged sample", "Valid" if ee["valid"] else "N/A", "Not claimed as task completion accuracy", "Raw_Data"),
        ("Overall", "EE mean 3D position error", ee["mean"], "mm", ee["valid"], "arithmetic mean", "Valid" if ee["valid"] else "N/A", "", "Raw_Data"),
        ("Overall", "EE population standard deviation", ee["std"], "mm", ee["valid"], "population std (ddof=0)", "Valid" if ee["valid"] else "N/A", "", "Raw_Data"),
        ("Moving samples", "EE 3D position RMSE", moving_ee["rmse"], "mm", moving_ee["valid"], "RMSE on activity-threshold samples", "Valid" if moving_ee["valid"] else "N/A", f"Threshold {calc['activity_threshold']:.1f} deg/s", "Raw_Data"),
        ("Moving samples", "EE maximum 3D position error", moving_ee["max"], "mm", moving_ee["valid"], "maximum on activity-threshold samples", "Valid" if moving_ee["valid"] else "N/A", "", "Raw_Data"),
        ("Availability", "Payload metadata", None, "kg", len(valid(by_name["payload_kg"])), "valid payload_kg observations", availability(by_name, ["payload_kg"]), "Missing is not zero", "Raw_Data"),
        ("Availability", "Commanded effort", None, "N·m", 0, "valid commanded effort observations", availability(by_name, ["__commanded_effort_Nm"]), "Missing is not zero", "Raw_Data"),
        ("Availability", "Measured effort", None, "N·m", 0, "valid measured effort observations", availability(by_name, ["__measured_effort_Nm"]), "Missing is not zero", "Raw_Data"),
        ("Availability", "Logged PID gains", None, "", 0, "valid Kp/Ki/Kd observations", availability(by_name, ["__Kp", "__Ki", "__Kd"]), "Absence does not imply no internal controller feedback", "Raw_Data"),
    ])
    start = 4
    for r, row in enumerate(rows, start):
        for c, value in enumerate(row, 1):
            set_cell(ws, r, c, value, border=Border(bottom=THIN), align=Alignment(vertical="top", wrap_text=True))
        if isinstance(row[2], float):
            ws.cell(r, 3).number_format = "0.000000"
        if r % 2 == 0:
            for cell in ws[r]: cell.fill = PatternFill("solid", fgColor="F7F9FC")
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:I{start+len(rows)-1}"
    widths = [18, 39, 16, 12, 15, 43, 23, 52, 18]
    for i, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = width
    ws.conditional_formatting.add(f"G{start}:G{start+len(rows)-1}", FormulaRule(formula=[f'LEFT(G{start},3)="Not"'], fill=PatternFill("solid", fgColor="FFF2CC")))
    return ws


def write_run_info(wb, input_path, source_hash, headers, by_name, calc, created):
    ws = wb.create_sheet("Run_Info")
    style_sheet(ws, False, 95)
    title_sheet(ws, "UR5 Trajectory Analysis — Run Information and Data Integrity", 4)
    for c, name in enumerate(("Field", "Value", "Meaning / method", "Status"), 1):
        set_cell(ws, 3, c, name, fill="4472C4", font=Font(name="Aptos", bold=True, color=WHITE), align=Alignment(horizontal="center"))
    invalid_time = len(by_name["time_sec"]) - len(valid(by_name["time_sec"]))
    rows = [
        ("Source CSV", str(input_path.resolve()), "Read-only source used to create this report", "Recorded"),
        ("Source SHA-256", source_hash, "Checksum at report generation", "Recorded"),
        ("Report created (UTC)", created.isoformat(), "Timestamp generated by this offline tool", "Recorded"),
        ("Trajectory ID", unique_text(by_name["trajectory_id"]), "Distinct nonblank logged values", "Recorded"),
        ("Controller mode label", unique_text(by_name["controller_mode"]), "Logged metadata; not an audit of internal controller implementation", "Recorded"),
        ("Payload", availability(by_name, ["payload_kg"]), "No physical payload is inferred", availability(by_name, ["payload_kg"])),
        ("Samples", len(by_name["time_sec"]), "CSV data rows", "Valid"),
        ("First timestamp", calc["first_time"], "Absolute logged time_sec", "Valid" if calc["first_time"] is not None else "N/A"),
        ("Last timestamp", calc["last_time"], "Absolute logged time_sec", "Valid" if calc["last_time"] is not None else "N/A"),
        ("Logged duration", calc["duration"], "last valid timestamp − first valid timestamp; not claimed as execution cycle time", "Valid" if calc["duration"] is not None else "N/A"),
        ("Median positive sample interval", calc["median_dt"], "Median of positive adjacent timestamp differences", "Valid" if calc["median_dt"] is not None else "N/A"),
        ("Invalid timestamp samples", invalid_time, "NaN/infinite/unparseable time_sec", "Valid" if invalid_time == 0 else "Review"),
        ("Non-positive timestamp intervals", calc["nonpositive_intervals"], "Repeated or backward adjacent valid timestamps", "Valid" if calc["nonpositive_intervals"] == 0 else "Review"),
        ("Raw units", "rad; rad/s; N·m; m; mm", "As encoded in CSV headers", "Recorded"),
        ("Joint error convention", "reference minus actual", "Logged error used only when it matches recomputation within 1e-9 rad", "Validated"),
        ("EE component error convention", "reference minus actual, signed", "(reference_m − actual_m) × 1000", "Checked"),
        ("EE 3D error", "Euclidean translational magnitude", "sqrt(x_error² + y_error² + z_error²), mm", "Checked"),
        ("Activity detection", f"max absolute actual joint velocity ≥ {calc['activity_threshold']:.1f} deg/s", "Moving-only metrics use threshold-satisfying samples; activity span is descriptive, not verified cycle bounds", "Inferred"),
        ("Activity span (relative)", f"{fmt(calc['activity_start'],3,' s')} to {fmt(calc['activity_end'],3,' s')}", "First/last threshold-satisfying samples", "Inferred" if calc["moving_count"] else "N/A"),
        ("Moving sample count", calc["moving_count"], "Samples meeting activity threshold", "Inferred"),
        ("Chart sampling", f"Full resolution ({len(by_name['time_sec']):,} rows)", "No chart-only downsampling", "Recorded"),
        ("EE orientation", "Not recorded", "Report assesses XYZ translation only, not 6-DOF pose", "Not recorded"),
        ("Commanded effort", availability(by_name, ["__commanded_effort_Nm"]), "Missing values remain blank/N/A", availability(by_name, ["__commanded_effort_Nm"])),
        ("Measured effort", availability(by_name, ["__measured_effort_Nm"]), "Missing values remain blank/N/A", availability(by_name, ["__measured_effort_Nm"])),
        ("PID gain metadata", availability(by_name, ["__Kp", "__Ki", "__Kd"]), "Absence does not imply no internal feedback", availability(by_name, ["__Kp", "__Ki", "__Kd"])),
        ("CSV columns", len(headers), "Original header order preserved in Raw_Data", "Valid"),
    ]
    r = 4
    for row in rows:
        for c, value in enumerate(row, 1):
            set_cell(ws, r, c, value, border=Border(bottom=THIN), align=Alignment(vertical="top", wrap_text=True))
        if r % 2 == 0:
            for cell in ws[r]: cell.fill = PatternFill("solid", fgColor="F7F9FC")
        r += 1
    r += 1
    merged_label(ws, f"A{r}:D{r}", "Joint error consistency checks", PALE_GREEN)
    r += 1
    for c, name in enumerate(("Joint", "Comparable samples", "Mismatches > 1e-9 rad", "Maximum absolute discrepancy (rad)"), 1):
        set_cell(ws, r, c, name, fill="70AD47", font=Font(bold=True, color=WHITE), align=Alignment(wrap_text=True))
    for jid, comparable, mismatches, maxdiff, use_logged in calc["joint_discrepancy"]:
        r += 1
        for c, value in enumerate((jid, comparable, mismatches, maxdiff), 1): set_cell(ws, r, c, value, border=Border(bottom=THIN))
        ws.cell(r, 4).number_format = "0.000000000000"
    r += 2
    merged_label(ws, f"A{r}:D{r}", "End-effector error consistency checks", PALE_ORANGE)
    r += 1
    for c, name in enumerate(("Signal", "Comparable samples", "Mismatches > 1e-6 mm", "Maximum absolute discrepancy (mm)"), 1):
        set_cell(ws, r, c, name, fill=ORANGE, font=Font(bold=True, color=WHITE), align=Alignment(wrap_text=True))
    for axis, comparable, mismatches, maxdiff in calc["ee_component_check"]:
        r += 1
        for c, value in enumerate((f"{axis} signed component", comparable, mismatches, maxdiff), 1): set_cell(ws, r, c, value, border=Border(bottom=THIN))
        ws.cell(r, 4).number_format = "0.000000000"
    comparable, mismatches, maxdiff = calc["ee_magnitude_check"]
    r += 1
    for c, value in enumerate(("3D magnitude", comparable, mismatches, maxdiff), 1): set_cell(ws, r, c, value, border=Border(bottom=THIN))
    ws.cell(r, 4).number_format = "0.000000000"
    ws.freeze_panes = "A4"
    for i, width in enumerate((31, 76, 74, 23), 1): ws.column_dimensions[get_column_letter(i)].width = width
    return ws


def write_comparisons(wb):
    ws = wb.create_sheet("Comparisons")
    style_sheet(ws, False, 100)
    title_sheet(ws, "UR5 Controller Comparisons — Reserved for Additional Logged Runs", 7)
    ws.merge_cells("A3:G6")
    set_cell(ws, 3, 1,
             "Awaiting additional logged runs. No baseline-vs-fixed-PID-vs-adaptive-PID values are fabricated here.\n\n"
             "Generate one workbook per run, then add a validated comparison process using identical metric definitions, payload conditions, trajectory, and logging boundaries.",
             fill=PALE_BLUE, font=Font(name="Aptos", size=14, bold=True, color=TEXT),
             align=Alignment(horizontal="center", vertical="center", wrap_text=True), border=Border(left=THIN,right=THIN,top=THIN,bottom=THIN))
    for col in range(1, 8): ws.column_dimensions[get_column_letter(col)].width = 22
    return ws


def scatter_chart(ws_data, title, y_title, series_defs, min_row, max_row, height=6.0, width=9.0, legend=True):
    chart = ScatterChart()
    chart.title = title
    chart.style = 13
    chart.height = height
    chart.width = width
    chart.x_axis.title = "Time (s)"
    chart.y_axis.title = y_title
    chart.x_axis.numFmt = "0"
    chart.x_axis.majorUnit = 20
    chart.y_axis.numFmt = "0.00"
    chart.x_axis.majorGridlines = ChartLines()
    chart.y_axis.majorGridlines = ChartLines()
    chart.legend.position = "t"
    chart.legend.overlay = False
    chart_text_style(chart, 1100)
    if not legend:
        chart.legend = None
    xvalues = Reference(ws_data, min_col=1, min_row=min_row, max_row=max_row)
    for col, name, color, dash in series_defs:
        yvalues = Reference(ws_data, min_col=col, min_row=min_row, max_row=max_row)
        series = Series(yvalues, xvalues, title=name)
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 16000
        if dash:
            series.graphicalProperties.line.prstDash = dash
        series.marker.symbol = "none"
        chart.series.append(series)
    return chart


def projection_chart(ws_data, title, x_col, y_defs, min_row, max_row, x_title, y_title, width=8.2):
    chart = ScatterChart()
    chart.title = title
    chart.style = 13
    chart.height = 5.8
    chart.width = width
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    chart.x_axis.numFmt = "0.00"
    chart.y_axis.numFmt = "0.00"
    chart.x_axis.majorGridlines = ChartLines()
    chart.y_axis.majorGridlines = ChartLines()
    chart_text_style(chart, 1100)
    chart.legend.position = "t"
    for xcol, ycol, name, color, dash in y_defs:
        xvalues = Reference(ws_data, min_col=xcol, min_row=min_row, max_row=max_row)
        yvalues = Reference(ws_data, min_col=ycol, min_row=min_row, max_row=max_row)
        series = Series(yvalues, xvalues, title=name)
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 16000
        if dash: series.graphicalProperties.line.prstDash = dash
        series.marker.symbol = "none"
        chart.series.append(series)
    return chart


def chart_text_style(chart, axis_size=1100):
    """Set explicit 11 pt chart text so Calc does not fall back to tiny labels."""
    axis_props = CharacterProperties(sz=axis_size, latin=DrawingFont(typeface="Liberation Sans"))
    title_props = CharacterProperties(sz=1300, b=True, latin=DrawingFont(typeface="Liberation Sans"))
    axis_text = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=axis_props))])
    title_text = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=title_props))])
    chart.x_axis.txPr = axis_text
    chart.y_axis.txPr = axis_text
    chart.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=title_props)
    if chart.legend is not None:
        chart.legend.txPr = axis_text


def velocity_chart(ws_data, title, velocity_col, error_col, min_row, max_row, height=7.4, width=12.8):
    """Scatter chart with time on the shared X axis and separate Y axes."""
    time_values = Reference(ws_data, min_col=1, min_row=min_row, max_row=max_row)
    chart = ScatterChart()
    chart.title = title
    chart.style = 13
    chart.height = height
    chart.width = width
    chart.x_axis.title = "Time (s)"
    chart.x_axis.numFmt = "0"
    chart.x_axis.majorUnit = 20
    chart.x_axis.majorGridlines = ChartLines()
    chart.y_axis.title = "Actual velocity (deg/s)"
    chart.y_axis.numFmt = "0"
    chart.y_axis.majorGridlines = ChartLines()
    chart.legend.position = "t"
    chart_text_style(chart, 1100)
    velocity_values = Reference(ws_data, min_col=velocity_col, min_row=min_row, max_row=max_row)
    velocity_series = Series(velocity_values, time_values, title="Actual velocity (deg/s)")
    velocity_series.graphicalProperties.line.solidFill = BLUE
    velocity_series.graphicalProperties.line.width = 15000
    velocity_series.marker.symbol = "none"
    chart.series.append(velocity_series)

    error_chart = ScatterChart()
    error_chart.y_axis.title = "Signed position error (deg)"
    error_chart.y_axis.numFmt = "0.00"
    error_chart.y_axis.axId = 200
    error_chart.y_axis.crosses = "max"
    error_chart.y_axis.majorGridlines = None
    error_chart.y_axis.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1100, latin=DrawingFont(typeface="Liberation Sans"))))])
    error_values = Reference(ws_data, min_col=error_col, min_row=min_row, max_row=max_row)
    error_series = Series(error_values, time_values, title="Signed position error (deg)")
    error_series.graphicalProperties.line.solidFill = RED
    error_series.graphicalProperties.line.width = 15000
    error_series.marker.symbol = "none"
    error_chart.series.append(error_series)
    chart += error_chart
    return chart


def analysis_header(ws, title, subtitle, last_col):
    style_sheet(ws, False, 85)
    title_sheet(ws, title, last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    set_cell(ws, 2, 1, subtitle, fill=PALE_BLUE, font=Font(name="Aptos", size=11, color=TEXT),
             align=Alignment(horizontal="left", vertical="center", wrap_text=True))
    ws.row_dimensions[2].height = 30
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def write_joint_tracking(wb, calc):
    ws = wb.create_sheet("Joint_Tracking")
    analysis_header(ws, "UR5 Joint Tracking — Enlarged Reference vs Actual", "All six joints use full-resolution calculated arrays. Reference is dashed blue; actual is solid red. Angles are degrees and time is relative seconds.", 30)
    for col in range(1, 31): ws.column_dimensions[get_column_letter(col)].width = 3.1
    data_ws = wb["Calculated_Data"]
    header_map = {name:i+1 for i, name in enumerate(calc["calc_headers"])}
    min_row, max_row = 4, len(calc["calc_rows"])+3
    for idx, (_, jid, label) in enumerate(JOINTS):
        row = 4 + (idx // 2) * 22
        anchor = f"A{row}" if idx % 2 == 0 else f"P{row}"
        chart = scatter_chart(data_ws, f"{jid} — {label}", "Position (deg)", [
            (header_map[f"{jid}_reference_deg"], "Reference", BLUE, "dash"),
            (header_map[f"{jid}_actual_deg"], "Actual", RED, None),
        ], min_row, max_row, height=7.8, width=13.3)
        ws.add_chart(chart, anchor)
    ws.merge_cells("A72:AD74")
    set_cell(ws, 72, 1, "Interpretation: the reference and actual trajectories are plotted at the original sample resolution; no interpolation or display-only resampling is used.", fill="E9EFF7", font=Font(size=10, color=TEXT), align=Alignment(wrap_text=True, vertical="center"))
    ws.print_area = "A1:AD74"
    return ws


def write_joint_errors(wb, calc):
    ws = wb.create_sheet("Joint_Errors")
    analysis_header(ws, "UR5 Joint Tracking Errors — Enlarged Signed Error Views", "Position error = reference minus actual. Green is signed error; dashed black is zero reference. Each title reports the absolute peak and relative timestamp.", 30)
    for col in range(1, 31): ws.column_dimensions[get_column_letter(col)].width = 3.1
    data_ws = wb["Calculated_Data"]
    header_map = {name:i+1 for i, name in enumerate(calc["calc_headers"])}
    min_row, max_row = 4, len(calc["calc_rows"])+3
    for idx, (_, jid, label) in enumerate(JOINTS):
        row = 4 + (idx // 2) * 22
        anchor = f"A{row}" if idx % 2 == 0 else f"P{row}"
        errors = calc["joint_arrays"][jid]["error"]
        signed, peak, peak_time = peak_abs_info(errors, calc["rel_time"])
        title = f"{jid} — {label} error | peak {peak:.3f} deg at t={peak_time:.2f} s" if peak is not None else f"{jid} — {label} error | N/A"
        chart = scatter_chart(data_ws, title, "Signed error (deg)", [
            (header_map[f"{jid}_error_deg"], "Position error", GREEN, None),
            (header_map["zero_reference"], "Zero", BLACK, "dash"),
        ], min_row, max_row, height=7.8, width=13.3)
        ws.add_chart(chart, anchor)
    ws.merge_cells("A72:AD74")
    set_cell(ws, 72, 1, "Peak annotations are computed from the full-resolution signed error arrays and use relative chart time. A peak is not a quality classification.", fill="FFF8E8", font=Font(size=10, color=TEXT), align=Alignment(wrap_text=True, vertical="center"))
    ws.print_area = "A1:AD74"
    return ws


def write_velocity_analysis(wb, calc):
    ws = wb.create_sheet("Velocity_Analysis")
    analysis_header(ws, "UR5 Joint Velocity vs Position Error — Enlarged Views", "Blue uses the left axis (actual velocity in deg/s); red uses the right axis (signed position error in deg). Time is relative seconds.", 30)
    for col in range(1, 31): ws.column_dimensions[get_column_letter(col)].width = 3.1
    data_ws = wb["Calculated_Data"]
    header_map = {name:i+1 for i, name in enumerate(calc["calc_headers"])}
    min_row, max_row = 4, len(calc["calc_rows"])+3
    for idx, (_, jid, label) in enumerate(JOINTS):
        row = 4 + (idx // 2) * 22
        anchor = f"A{row}" if idx % 2 == 0 else f"P{row}"
        errors = calc["joint_arrays"][jid]["error"]
        _, peak, peak_time = peak_abs_info(errors, calc["rel_time"])
        title = f"{jid} — {label} | error peak {peak:.3f} deg at t={peak_time:.2f} s" if peak is not None else f"{jid} — {label}"
        chart = velocity_chart(data_ws, title, header_map[f"{jid}_actual_velocity_deg_s"], header_map[f"{jid}_error_deg"], min_row, max_row)
        ws.add_chart(chart, anchor)
    ws.merge_cells("A72:AD74")
    set_cell(ws, 72, 1, "The two Y axes are intentionally independent: do not compare their vertical distances directly. All velocity and error calculations use the logged values after unit conversion.", fill="EDE2FA", font=Font(size=10, color=TEXT), align=Alignment(wrap_text=True, vertical="center"))
    ws.print_area = "A1:AD74"
    return ws


def write_ee_analysis(wb, by_name, calc, image_path):
    ws = wb.create_sheet("EE_Analysis")
    analysis_header(ws, "UR5 End-Effector Analysis — Trajectory and Cartesian Error", "Translation-only analysis in the base_link/tool0 frame. Reference is dashed blue, actual is solid red, signed components use X/Y/Z colors, and total Cartesian error is black.", 34)
    for col in range(1, 35): ws.column_dimensions[get_column_letter(col)].width = 3.0
    ws.column_dimensions["AG"].width = 20
    ws.column_dimensions["AH"].width = 14
    data_ws = wb["Calculated_Data"]
    header_map = {name:i+1 for i, name in enumerate(calc["calc_headers"])}
    min_row, max_row = 4, len(calc["calc_rows"])+3
    img = Image(str(image_path)); img.width = 760; img.height = 500; ws.add_image(img, "A4")
    xy = projection_chart(data_ws, "Top View — XY Plane", None, [
        (header_map["reference_ee_x_m"], header_map["reference_ee_y_m"], "Reference", BLUE, "dash"),
        (header_map["actual_ee_x_m"], header_map["actual_ee_y_m"], "Actual", RED, None),
    ], min_row, max_row, "X (m)", "Y (m)", width=12.0)
    xz = projection_chart(data_ws, "Side View — XZ Plane", None, [
        (header_map["reference_ee_x_m"], header_map["reference_ee_z_m"], "Reference", BLUE, "dash"),
        (header_map["actual_ee_x_m"], header_map["actual_ee_z_m"], "Actual", RED, None),
    ], min_row, max_row, "X (m)", "Z (m)", width=12.0)
    ws.add_chart(xy, "U4")
    ws.add_chart(xz, "AG4")
    xyz = scatter_chart(data_ws, "Signed XYZ Components and 3D Magnitude", "Position error (mm)", [
        (header_map["ee_error_x_mm"], "X signed", BLUE, None),
        (header_map["ee_error_y_mm"], "Y signed", RED, None),
        (header_map["ee_error_z_mm"], "Z signed", GREEN, None),
        (header_map["ee_error_3d_mm"], "3D magnitude", BLACK, None),
    ], min_row, max_row, height=8.0, width=16.0)
    ws.add_chart(xyz, "A31")
    ee = calc["overall_ee"]
    total = scatter_chart(data_ws, f"Total Cartesian Error — peak {fmt(ee['max'],3,' mm')} at t={fmt(ee['peak_time'],2,' s')}", "Total error (mm)", [
        (header_map["ee_error_3d_mm"], "3D magnitude", BLACK, None),
    ], min_row, max_row, height=8.0, width=16.0)
    ws.add_chart(total, "Q31")
    ws.merge_cells("AG31:AH31")
    set_cell(ws, 31, 33, "RMSE statistics", fill=PALE_PURPLE, font=Font(bold=True, color=TEXT), align=Alignment(horizontal="center"))
    stats = [("3D RMSE", ee["rmse"], "mm"), ("Maximum", ee["max"], "mm"), ("Final logged", ee["final"], "mm"), ("Mean", ee["mean"], "mm"), ("Population std", ee["std"], "mm")]
    for i, (name, value, unit) in enumerate(stats, 32):
        set_cell(ws, i, 33, name, fill="F7F9FC" if i % 2 == 0 else WHITE, border=Border(bottom=THIN))
        set_cell(ws, i, 34, value, fill=PALE_RED if name == "Maximum" else PALE_GREEN if name in ("3D RMSE", "Final logged") else WHITE, font=Font(bold=name in ("3D RMSE", "Maximum", "Final logged")), align=Alignment(horizontal="center"), number_format="0.000")
    ws.merge_cells("AG39:AH42")
    set_cell(ws, 39, 33, "Final logged means the last valid logged sample; it is not claimed as verified task completion accuracy. No EE orientation signal is present.", fill="E9EFF7", font=Font(size=10, color=TEXT), align=Alignment(wrap_text=True, vertical="top"))
    ws.merge_cells("A54:AH56")
    set_cell(ws, 54, 1, "The 3D trajectory is a static embedded plot for LibreOffice portability. Native Excel/Calc charts show the time-series and XY/XZ projections. All metrics use the full-resolution source data.", fill="FFF8E8", font=Font(size=10, color=TEXT), align=Alignment(wrap_text=True, vertical="center"))
    ws.print_area = "A1:AH56"
    return ws


def create_3d_plot(path, by_name):
    fig = plt.figure(figsize=(8.2, 6.1), dpi=180)
    ax = fig.add_subplot(111, projection="3d")
    rr = [by_name[f"reference_ee_{axis}_m"] for axis in "xyz"]
    aa = [by_name[f"actual_ee_{axis}_m"] for axis in "xyz"]
    valid_ref = [(x,y,z) for x,y,z in zip(*rr) if finite(x) and finite(y) and finite(z)]
    valid_act = [(x,y,z) for x,y,z in zip(*aa) if finite(x) and finite(y) and finite(z)]
    if valid_ref:
        x,y,z = zip(*valid_ref); ax.plot(x,y,z, color="#1F5BFF", linestyle="--", linewidth=1.5, label="Reference")
    if valid_act:
        x,y,z = zip(*valid_act); ax.plot(x,y,z, color="#E31A1C", linewidth=1.4, label="Actual")
    ax.set_xlabel("X (m)"); ax.set_ylabel("Y (m)"); ax.set_zlabel("Z (m)")
    ax.grid(True, alpha=.25); ax.legend(loc="upper right")
    ax.view_init(elev=25, azim=-55)
    fig.tight_layout()
    fig.savefig(path, transparent=False, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def dashboard_panel(ws, start_col, end_col, row, title, fill, color=TEXT):
    merged_label(ws, f"{get_column_letter(start_col)}{row}:{get_column_letter(end_col)}{row}", title, fill, 11, color)
    ws.row_dimensions[row].height = 22


def write_dashboard(wb, by_name, calc, image_path):
    ws = wb.create_sheet("Dashboard", 0)
    style_sheet(ws, False, 65)
    ws.freeze_panes = "A6"
    for c in range(1, 81):
        ws.column_dimensions[get_column_letter(c)].width = 2.35
    for r in range(1, 120): ws.row_dimensions[r].height = 16
    ws.merge_cells("A1:CB2")
    set_cell(ws, 1, 1, "UR5 Pick-and-Place — Trajectory Tracking Analysis", fill=NAVY,
             font=Font(name="Aptos Display", size=22, bold=True, color=WHITE),
             align=Alignment(horizontal="left", vertical="center"))
    trajectory = unique_text(by_name["trajectory_id"])
    controller = unique_text(by_name["controller_mode"])
    ws.merge_cells("A3:CB3")
    set_cell(ws, 3, 1, f"Run: {trajectory}  |  Controller-mode label: {controller}  |  Offline evidence report",
             fill="274F7D", font=Font(name="Aptos", size=11, bold=True, color=WHITE), align=Alignment(vertical="center"))
    kpis = [
        ("Total Samples", f"{len(by_name['time_sec']):,}", PALE_BLUE, NAVY),
        ("Logged Duration", fmt(calc["duration"],2," s"), PALE_BLUE, NAVY),
        ("Median Sample Interval", fmt(calc["median_dt"]*1000 if calc["median_dt"] is not None else None,1," ms"), PALE_BLUE, NAVY),
        ("EE Position RMSE", fmt(calc["overall_ee"]["rmse"],2," mm"), PALE_GREEN, "087A25"),
        ("EE Max Error", fmt(calc["overall_ee"]["max"],2," mm"), PALE_RED, "C00000"),
        ("Final Logged EE Error", fmt(calc["overall_ee"]["final"],2," mm"), PALE_GREEN, "087A25"),
        ("Controller Mode", controller, PALE_BLUE, NAVY),
    ]
    widths = [(1,11),(13,23),(25,35),(37,47),(49,59),(61,70),(72,80)]
    for (label, value, fill, color), (c1,c2) in zip(kpis, widths):
        ws.merge_cells(start_row=4,start_column=c1,end_row=4,end_column=c2)
        ws.merge_cells(start_row=5,start_column=c1,end_row=5,end_column=c2)
        set_cell(ws,4,c1,label,fill=fill,font=Font(name="Aptos",size=10,bold=True,color=TEXT),align=Alignment(horizontal="center",vertical="center"),border=Border(left=THIN,right=THIN,top=THIN))
        set_cell(ws,5,c1,value,fill=fill,font=Font(name="Aptos Display",size=14,bold=True,color=color),align=Alignment(horizontal="center",vertical="center",wrap_text=True),border=Border(left=THIN,right=THIN,bottom=THIN))
    dashboard_panel(ws,1,36,7,"1. Joint Position Tracking (Reference vs Actual)",PALE_BLUE,NAVY)
    dashboard_panel(ws,37,64,7,"2. Joint Position Tracking Error",PALE_GREEN,"145A14")
    dashboard_panel(ws,65,80,7,"3. Joint Error Statistics",PALE_ORANGE,"7F4A00")
    dashboard_panel(ws,65,80,31,"4. End-Effector Error Statistics",PALE_PURPLE,"312070")
    dashboard_panel(ws,1,24,62,"5. End-Effector Trajectory (3D View)",PALE_ORANGE,"8B2500")
    dashboard_panel(ws,25,52,62,"6. End-Effector Position Error (XYZ)",PALE_BLUE,NAVY)
    dashboard_panel(ws,53,80,62,"7. End-Effector Total Error",PALE_RED,"7A1730")
    dashboard_panel(ws,25,52,82,"8. Joint Velocity vs Position Error — J5",PALE_PURPLE,"312070")
    dashboard_panel(ws,53,80,82,"9. XY and XZ End-Effector Projection",PALE_GREEN,"145A14")

    data_ws = wb["Calculated_Data"]
    header_map = {name:i+1 for i,name in enumerate(calc["calc_headers"])}
    min_row, max_row = 4, len(calc["calc_rows"])+3
    track_anchors = ["A8","S8","A26","S26","A44","S44"]
    error_anchors = ["AK8","AY8","AK26","AY26","AK44","AY44"]
    for (_, jid, label), t_anchor, e_anchor in zip(JOINTS, track_anchors, error_anchors):
        ref_col=header_map[f"{jid}_reference_deg"]; act_col=header_map[f"{jid}_actual_deg"]; err_col=header_map[f"{jid}_error_deg"]
        chart=scatter_chart(data_ws,f"{jid} ({label})","Position (deg)",[(ref_col,"Reference",BLUE,"dash"),(act_col,"Actual",RED,None)],min_row,max_row,5.8,10.0)
        ws.add_chart(chart,t_anchor)
        err_chart=scatter_chart(data_ws,f"{jid} Position Error","Error (deg)",[(err_col,"Signed error",GREEN,None),(header_map['zero_reference'],"Zero",BLACK,"dash")],min_row,max_row,5.8,8.0,False)
        ws.add_chart(err_chart,e_anchor)

    table_row=9
    table_ranges=((65,70,"Joint"),(71,75,"RMSE (deg)"),(76,80,"Max abs (deg)"))
    for start_col, end_col, name in table_ranges:
        ws.merge_cells(start_row=table_row, start_column=start_col, end_row=table_row, end_column=end_col)
        set_cell(ws,table_row,start_col,name,fill="D9E2F3",font=Font(bold=True,color=TEXT),align=Alignment(horizontal="center",wrap_text=True),border=Border(left=THIN,right=THIN,top=THIN,bottom=THIN))
    joint_stats=[]
    for _,jid,_ in JOINTS:
        e=calc["joint_arrays"][jid]["error"]; joint_stats.append((jid,rmse(e),max_abs(e)))
    max_rmse=max(v for _,v,_ in joint_stats if v is not None); max_err=max(v for _,_,v in joint_stats if v is not None)
    for r,(jid,rms,mx) in enumerate(joint_stats,table_row+1):
        values=((65,70,jid),(71,75,rms),(76,80,mx))
        for start_col, end_col, value in values:
            ws.merge_cells(start_row=r, start_column=start_col, end_row=r, end_column=end_col)
            is_peak = (start_col == 71 and value == max_rmse) or (start_col == 76 and value == max_err)
            fill="FCE4D6" if is_peak else ("F7F9FC" if r%2==0 else WHITE)
            set_cell(ws,r,start_col,value,fill=fill,font=Font(bold=(start_col==65 or is_peak)),align=Alignment(horizontal="center"),border=Border(left=THIN,right=THIN,top=THIN,bottom=THIN),number_format="0.000")
    ws.merge_cells(start_row=18,start_column=65,end_row=19,end_column=80)
    set_cell(ws,18,65,"Highlight = largest observed joint metric; no quality classification is implied.",fill="FFF8E8",font=Font(size=9,italic=True,color=TEXT),align=Alignment(wrap_text=True,vertical="center"))

    ee=calc["overall_ee"]
    ee_rows=[("3D RMSE (mm)",ee["rmse"]),("Maximum (mm)",ee["max"]),("Final logged sample (mm)",ee["final"]),("Mean (mm)",ee["mean"]),("Population std (mm)",ee["std"])]
    for start_col, end_col, name in ((65,73,"Metric"),(74,80,"Value")):
        ws.merge_cells(start_row=33,start_column=start_col,end_row=33,end_column=end_col)
        set_cell(ws,33,start_col,name,fill="D9E2F3",font=Font(bold=True),align=Alignment(horizontal="center"),border=Border(left=THIN,right=THIN,top=THIN,bottom=THIN))
    for r,(name,value) in enumerate(ee_rows,34):
        ws.merge_cells(start_row=r,start_column=65,end_row=r,end_column=73)
        ws.merge_cells(start_row=r,start_column=74,end_row=r,end_column=80)
        set_cell(ws,r,65,name,fill=("F7F9FC" if r%2==0 else WHITE),border=Border(left=THIN,right=THIN,top=THIN,bottom=THIN))
        color="C00000" if name.startswith("Maximum") else ("087A25" if name.startswith("3D") or name.startswith("Final") else TEXT)
        fill=PALE_RED if name.startswith("Maximum") else (PALE_GREEN if name.startswith("3D") or name.startswith("Final") else ("F7F9FC" if r%2==0 else WHITE))
        set_cell(ws,r,74,value,fill=fill,font=Font(bold=name.startswith(("3D","Maximum","Final")),color=color),align=Alignment(horizontal="center"),border=Border(left=THIN,right=THIN,top=THIN,bottom=THIN),number_format="0.000")
    ws.merge_cells(start_row=41,start_column=65,end_row=43,end_column=80)
    set_cell(ws,41,65,f"Peak: {fmt(ee['max'],3,' mm')} at t = {fmt(ee['peak_time'],3,' s')}\nFinal means last valid logged sample, not verified task completion.\nTranslational XYZ only; orientation not recorded.",fill="F7F3FC",font=Font(size=9,color=TEXT),align=Alignment(wrap_text=True,vertical="top"))

    img=Image(str(image_path)); img.width=610; img.height=410; ws.add_image(img,"A64")
    ws.merge_cells("A90:X92")
    set_cell(ws,90,1,"Static embedded matplotlib view generated from full-resolution CSV data; it is not Excel-rotatable.",fill="FFF8E8",font=Font(size=9,italic=True,color=TEXT),align=Alignment(wrap_text=True,vertical="center"))

    ee_chart=scatter_chart(data_ws,"Signed XYZ Components and 3D Magnitude","Position error (mm)",[
        (header_map['ee_error_x_mm'],"X signed",BLUE,None),(header_map['ee_error_y_mm'],"Y signed",RED,None),(header_map['ee_error_z_mm'],"Z signed",GREEN,None),(header_map['ee_error_3d_mm'],"3D magnitude",BLACK,None)],min_row,max_row,6.8,14.0)
    ws.add_chart(ee_chart,"Y64")
    total_chart=scatter_chart(data_ws,f"3D Error — peak {fmt(ee['max'],2,' mm')} at {fmt(ee['peak_time'],2,' s')}","Total error (mm)",[(header_map['ee_error_3d_mm'],"3D magnitude",BLACK,None)],min_row,max_row,6.8,14.0)
    ws.add_chart(total_chart,"BA64")

    vel_col=header_map['J5_actual_velocity_deg_s']; err_col=header_map['J5_error_deg']
    time_values=Reference(data_ws,min_col=1,min_row=min_row,max_row=max_row)
    chart=ScatterChart(); chart.title="J5 Actual Velocity vs Signed Position Error"; chart.style=13; chart.height=6.8; chart.width=14.0
    chart.x_axis.title="Time (s)"; chart.y_axis.title="Velocity (deg/s)"; chart.legend.position="t"
    vel_values=Reference(data_ws,min_col=vel_col,min_row=min_row,max_row=max_row)
    vel_series=Series(vel_values,time_values,title="J5 actual velocity")
    vel_series.graphicalProperties.line.solidFill=BLUE; vel_series.graphicalProperties.line.width=13000; vel_series.marker.symbol="none"
    chart.series.append(vel_series)
    chart2=ScatterChart(); err_values=Reference(data_ws,min_col=err_col,min_row=min_row,max_row=max_row)
    err_series=Series(err_values,time_values,title="J5 signed position error")
    err_series.graphicalProperties.line.solidFill=RED; err_series.graphicalProperties.line.width=13000; err_series.marker.symbol="none"
    chart2.series.append(err_series)
    chart2.y_axis.title="Signed error (deg)"; chart2.y_axis.axId=200; chart2.y_axis.crosses="max"
    chart += chart2; ws.add_chart(chart,"Y84")

    xy=projection_chart(data_ws,"Top View — XY Plane",None,[
        (header_map['reference_ee_x_m'],header_map['reference_ee_y_m'],"Reference",BLUE,"dash"),(header_map['actual_ee_x_m'],header_map['actual_ee_y_m'],"Actual",RED,None)],min_row,max_row,"X (m)","Y (m)",6.2)
    xz=projection_chart(data_ws,"Side View — XZ Plane",None,[
        (header_map['reference_ee_x_m'],header_map['reference_ee_z_m'],"Reference",BLUE,"dash"),(header_map['actual_ee_x_m'],header_map['actual_ee_z_m'],"Actual",RED,None)],min_row,max_row,"X (m)","Z (m)",6.2)
    ws.add_chart(xy,"BA84"); ws.add_chart(xz,"BM84")

    ws.merge_cells("A113:CB116")
    set_cell(ws,113,1,
             "Methods/limitations: metrics use every valid source sample; no interpolation or chart downsampling. Logged duration is the timestamp span, not a verified task-cycle duration. Error convention is reference minus actual. Missing payload, effort, and gains remain N/A. The report covers end-effector XYZ translation only; no measured orientation is present. See Metrics and Run_Info for validation details.",
             fill="E9EFF7",font=Font(name="Aptos",size=9,color=TEXT),align=Alignment(wrap_text=True,vertical="top"),border=Border(top=THIN,bottom=THIN))
    ws.print_area="A1:CB116"
    ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
    return ws


def build_workbook(input_path: Path, output_path: Path, overwrite: bool):
    if input_path.resolve() == output_path.resolve():
        raise ValueError("Input CSV and output XLSX must be different files")
    if input_path.suffix.lower() != ".csv":
        raise ValueError("Input must be a .csv file")
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("Output must end in .xlsx")
    if not input_path.is_file():
        raise FileNotFoundError(f"Input does not exist: {input_path}")
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"Output exists: {output_path}. Use --overwrite to replace it explicitly.")
    backup_path = None
    if output_path.exists() and overwrite:
        stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        backup_path = output_path.with_name(f"{output_path.stem}.backup_{stamp}{output_path.suffix}")
        shutil.copy2(output_path, backup_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    before_hash = sha256(input_path)
    before_size = input_path.stat().st_size
    headers, raw_rows, values, by_name = load_csv(input_path)
    calc = calculate(by_name, len(raw_rows))
    created = datetime.now(timezone.utc)
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    write_raw_data(wb, headers, raw_rows)
    write_calculated(wb, calc)
    write_metrics(wb, by_name, calc)
    write_run_info(wb, input_path, before_hash, headers, by_name, calc, created)
    write_comparisons(wb)
    with tempfile.TemporaryDirectory(prefix="ur5_excel_") as temp_dir:
        image_path = Path(temp_dir) / "ee_trajectory_3d.png"
        create_3d_plot(image_path, by_name)
        write_dashboard(wb, by_name, calc, image_path)
        write_joint_tracking(wb, calc)
        write_joint_errors(wb, calc)
        write_ee_analysis(wb, by_name, calc, image_path)
        write_velocity_analysis(wb, calc)
        sheet_order = ("Dashboard", "Joint_Tracking", "Joint_Errors", "EE_Analysis", "Velocity_Analysis",
                       "Raw_Data", "Calculated_Data", "Metrics", "Run_Info", "Comparisons")
        wb._sheets = [wb[name] for name in sheet_order]
        temp_output = output_path.parent / f".{output_path.name}.{os.getpid()}.tmp.xlsx"
        try:
            wb.save(temp_output)
            check = load_workbook(temp_output, read_only=False, data_only=False)
            expected = ["Dashboard", "Joint_Tracking", "Joint_Errors", "EE_Analysis", "Velocity_Analysis",
                        "Raw_Data", "Calculated_Data", "Metrics", "Run_Info", "Comparisons"]
            if check.sheetnames != expected:
                raise RuntimeError(f"Workbook verification failed: sheets are {check.sheetnames}")
            if len(check["Dashboard"]._charts) != 17:
                raise RuntimeError(f"Workbook verification failed: expected 17 native dashboard charts, found {len(check['Dashboard']._charts)}")
            expected_chart_counts = {"Joint_Tracking": 6, "Joint_Errors": 6, "EE_Analysis": 4, "Velocity_Analysis": 6}
            for sheet_name, expected_count in expected_chart_counts.items():
                actual_count = len(check[sheet_name]._charts)
                if actual_count != expected_count:
                    raise RuntimeError(f"Workbook verification failed: {sheet_name} expected {expected_count} charts, found {actual_count}")
            check.close()
            os.replace(temp_output, output_path)
        finally:
            if temp_output.exists(): temp_output.unlink()
    after_hash = sha256(input_path)
    after_size = input_path.stat().st_size
    if before_hash != after_hash or before_size != after_size:
        raise RuntimeError("Source CSV changed during report generation; output was created but source integrity check failed")
    return {
        "samples": len(raw_rows), "columns": len(headers), "duration": calc["duration"],
        "median_dt": calc["median_dt"], "ee": calc["overall_ee"], "sha256": before_hash,
        "size": before_size, "charts": 17, "backup": str(backup_path) if backup_path else None,
    }


def main() -> int:
    args = parse_args()
    try:
        result = build_workbook(args.input, args.output, args.overwrite)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(f"Created: {args.output.resolve()}")
    print(f"Source preserved: {args.input.resolve()} ({result['size']} bytes, SHA-256 {result['sha256']})")
    print(f"Samples: {result['samples']:,}; columns: {result['columns']}; logged duration: {fmt(result['duration'],3,' s')}; median interval: {fmt(result['median_dt'],6,' s')}")
    print(f"EE 3D RMSE: {fmt(result['ee']['rmse'],6,' mm')}; max: {fmt(result['ee']['max'],6,' mm')}; final logged: {fmt(result['ee']['final'],6,' mm')}")
    print("Verified sheets: Dashboard, Joint_Tracking, Joint_Errors, EE_Analysis, Velocity_Analysis, Raw_Data, Calculated_Data, Metrics, Run_Info, Comparisons")
    print(f"Verified native dashboard charts: {result['charts']} plus enlarged analysis charts (and embedded static 3D plots)")
    if result["backup"]:
        print(f"Backup saved: {result['backup']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
