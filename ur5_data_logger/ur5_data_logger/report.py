import csv
import json
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .schema import CSV_COLUMNS, EE_COLUMNS, JOINTS, METRIC_COLUMNS


NAVY = "17365D"
WHITE = "FFFFFF"
PASTELS = ("D9EAF7", "E2F0D9", "FFF2CC", "FCE4D6", "E4DFEC", "DDEBF7", "F4CCCC", "D9D2E9", "CFE2F3", "FCE5CD")
THIN_GRAY = Side(style="thin", color="D9E1F2")


def write_summary(path: Path, rows):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=METRIC_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def _title(ws, text, last_column):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    cell = ws.cell(1, 1, text)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True, size=16)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28


def _style_table(ws, header_row, max_row, max_column):
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max_column)}{max_row}"
    for cell in ws[header_row]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in range(header_row + 1, max_row + 1):
        if (row - header_row) % 2 == 0:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor="F5F8FC")
        for cell in ws[row]:
            cell.border = Border(bottom=THIN_GRAY)
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"
    ws.row_dimensions[header_row].height = 42


def _tracking_sheet(workbook, records):
    ws = workbook.active
    ws.title = "Joint Tracking"
    _title(ws, "UR5 Experiment — Joint Tracking", len(CSV_COLUMNS))
    groups = [("Run Metadata", 4)] + [(field.replace("_", " ").title(), len(JOINTS)) for field in (
        "reference_position_rad", "actual_position_rad", "position_error_rad",
        "reference_velocity_rad_s", "actual_velocity_rad_s", "commanded_effort_Nm",
        "measured_effort_Nm", "Kp", "Ki", "Kd")]
    groups.append(("End-Effector Tracking", len(EE_COLUMNS)))
    column = 1
    for index, (name, width) in enumerate(groups):
        ws.merge_cells(start_row=2, start_column=column, end_row=2, end_column=column + width - 1)
        cell = ws.cell(2, column, name)
        cell.fill = PatternFill("solid", fgColor=PASTELS[index % len(PASTELS)])
        cell.font = Font(bold=True, color="1F1F1F")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        column += width
    for col, name in enumerate(CSV_COLUMNS, 1):
        ws.cell(3, col, name)
    for row_index, record in enumerate(records, 4):
        for col, name in enumerate(CSV_COLUMNS, 1):
            value = record.get(name)
            ws.cell(row_index, col, None if value is None or (isinstance(value, float) and math.isnan(value)) else value)
    _style_table(ws, 3, max(3, ws.max_row), len(CSV_COLUMNS))
    for col, name in enumerate(CSV_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col)].width = 17 if col > 4 else (20 if col in (2, 3) else 14)
        if "position_error_rad" in name and ws.max_row >= 4:
            letter = get_column_letter(col)
            ws.conditional_formatting.add(
                f"{letter}4:{letter}{ws.max_row}",
                ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0,
                               mid_color="63BE7B", end_type="max", end_color="F8696B"),
            )
    return ws


def _simple_sheet(workbook, title, heading, columns, rows):
    ws = workbook.create_sheet(title)
    _title(ws, heading, len(columns))
    for col, name in enumerate(columns, 1):
        ws.cell(2, col, name)
    for row_index, record in enumerate(rows, 3):
        for col, name in enumerate(columns, 1):
            value = record.get(name)
            ws.cell(row_index, col, None if value is None or (isinstance(value, float) and math.isnan(value)) else value)
    _style_table(ws, 2, max(2, ws.max_row), len(columns))
    for col, name in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(col)].width = max(15, min(42, len(name) + 4))
    return ws


def write_workbook(path: Path, records, metrics, info, ee_verified):
    workbook = Workbook()
    _tracking_sheet(workbook, records)
    _simple_sheet(workbook, "Summary Metrics", "UR5 Experiment — Summary Metrics", METRIC_COLUMNS, metrics)
    info_rows = []
    for key, value in info.items():
        if isinstance(value, (dict, list, tuple)):
            value = json.dumps(value, sort_keys=True)
        info_rows.append({"field": key, "value": value})
    _simple_sheet(workbook, "Experiment Info", "UR5 Experiment — Provenance", ("field", "value"), info_rows)
    if ee_verified:
        ee_rows = [{"time_sec": record["time_sec"], **{key: record[key] for key in EE_COLUMNS}} for record in records]
        _simple_sheet(workbook, "EE Tracking", "UR5 Experiment — End-Effector Tracking", ("time_sec",) + EE_COLUMNS, ee_rows)
    workbook.save(path)
